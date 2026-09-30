"""Read a plan workbook into the inputs model.py needs.

The workbook is the source of truth. Its formulas have to be calculated before their values can be read, so the
file is first recalculated with LibreOffice (headless) into a temporary copy. Without LibreOffice, the values
Excel / LibreOffice / Numbers saved in the file are used; a file just written by these scripts has none, so open
and save it once.
"""
import copy
import shutil
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import tax_tables as TX
from build_workbook import BUCKETS, N_EVENTS, N_KIDS, N_PROPS, PROP_COLS, RESULTS, RESULTS_ROW0, Book, model_layout


def soffice():
    for name in ("soffice", "libreoffice"):
        if shutil.which(name):
            return shutil.which(name)
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def recalculated(path):
    """Path to a copy of the workbook with every formula calculated (or the file itself if LibreOffice is missing)."""
    exe = soffice()
    if not exe:
        return Path(path), False
    out, profile = Path(tempfile.mkdtemp()), Path(tempfile.mkdtemp())
    subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--calc", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(path)], check=True, capture_output=True)
    shutil.rmtree(profile, ignore_errors=True)
    return out / Path(path).name, True


def num(v, default=0.0):
    if v is None or v == "":
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def year(v):
    return None if v in (None, "") else int(float(v))


def kv(ws, col_key=1, col_val=2, row0=2):
    out = {}
    for r in range(row0, ws.max_row + 1):
        k = ws.cell(r, col_key).value
        if k in (None, ""):
            break
        out[str(k)] = ws.cell(r, col_val).value
    return out


def column_keys(country):
    """Year-table column keys in order (the same list the builder writes)."""
    b = Book.__new__(Book)
    b.country, b.us = country, country == "US"
    b.I = b.H = b.HC = b.TT = b.BR = b.pc = defaultdict(str)
    b.prop_rows = b.ev_rows = (0, 0)
    b.budget_rng = ("", "")
    return [col[0] for col in b.model_columns()]


def column_letters(country):
    return {k: get_column_letter(i + 1) for i, k in enumerate(column_keys(country))}


def tax_positions(country):
    """(col, first_row, last_row) of each table on the tax_tables tab, as the builder lays them out."""
    pos, col = {}, 5
    brackets = TX.US_BRACKETS if country == "US" else TX.CA_BRACKETS
    for key in brackets:
        pos[key] = (col, 5, 4 + TX.BRACKET_ROWS)
        col += 4
    if country == "US":
        pos["irmaa"] = (col, 5, 4 + len(TX.US_IRMAA))
        col += 5
        pos["salt"] = (col, 5, 4 + len(TX.US_SALT))
        col += 3
        pos["factors"] = (col, 5, 4 + len(TX.UNIFORM_LIFETIME))
    else:
        pos["factors"] = (col, 5, 4 + len(TX.RRIF_FACTORS))
    return pos


def read_tax(ws, country):
    T = {k: (v if isinstance(v, str) else num(v)) for k, v in kv(ws, row0=4).items()}
    T["TAX_YEAR"] = int(T["TAX_YEAR"])
    pos = tax_positions(country)
    for key, (col, a, b) in pos.items():
        rows = [[ws.cell(r, col + j).value for j in range(4)] for r in range(a, b + 1)]
        if key in ("fed", "ltcg", "state", "prov"):
            T[key] = [(num(r[0]), num(r[1])) for r in rows if r[1] not in (None, "")]
        elif key == "irmaa":
            T[key] = [(num(r[0]), num(r[1]), num(r[2])) for r in rows if r[0] not in (None, "")]
        elif key == "salt":
            T[key] = [(int(r[0]), num(r[1])) for r in rows if r[0] not in (None, "")]
        else:
            T["uniform" if country == "US" else "rrif"] = {int(r[0]): num(r[1]) for r in rows if r[0] not in (None, "")}
    return T


def read_shared(wb):
    hh_ws = wb["household"]
    hh = kv(hh_ws)
    country = str(hh["country"]).strip().upper()
    kids = []
    for r in range(1, hh_ws.max_row + 1):
        if hh_ws.cell(r, 1).value == "name" and hh_ws.cell(r, 2).value == "born":
            for rr in range(r + 1, r + 1 + N_KIDS):
                if hh_ws.cell(rr, 1).value not in (None, ""):
                    kids.append(dict(name=hh_ws.cell(rr, 1).value, born=year(hh_ws.cell(rr, 2).value),
                                     health_through=year(hh_ws.cell(rr, 3).value) if country == "US" else None))
            break
    people = []
    for k in (1, 2):
        p = dict(name=hh[f"p{k}_name"], born=int(num(hh[f"p{k}_born"])),
                 sex=0 if str(hh.get(f"p{k}_sex") or "F").strip().upper().startswith("M") else 1)
        if country == "US":
            p["benefit"] = num(hh.get(f"p{k}_ss_benefit"))
        else:
            p["benefit"] = num(hh.get(f"p{k}_cpp_at_65"))
            p["oas_years"] = num(hh.get(f"p{k}_oas_years"), 40.0)
        people.append(p)
    ws = wb["budget"]
    household = sum(num(ws.cell(r, 6).value) for r in range(2, 301) if ws.cell(r, 1).value == "household")
    ws = wb["kids_budget"]
    kids_lines = [dict(kid=ws.cell(r, 1).value, amount=num(ws.cell(r, 3).value), start=num(ws.cell(r, 4).value),
                       end=num(ws.cell(r, 5).value), inflation=ws.cell(r, 6).value, paid_from=ws.cell(r, 7).value)
                  for r in range(2, 41)]
    ws = wb["other_budget"]
    other_lines = [dict(amount=num(ws.cell(r, 2).value), start=num(ws.cell(r, 3).value), end=num(ws.cell(r, 4).value))
                   for r in range(2, 31)]
    ws = wb["properties"]
    head = [ws.cell(1, j).value for j in range(1, len(PROP_COLS) + 1)]
    assert head == PROP_COLS, "properties tab: the header row has changed"
    props = []
    for r in range(2, 2 + N_PROPS):
        v = {h: ws.cell(r, j + 1).value for j, h in enumerate(head)}
        p = {k: num(v[k]) for k in PROP_COLS if k not in ("property", "type", "notes", "depreciation_end_year")}
        p["name"] = v["property"] or ""
        p["depreciation_end_year"] = year(v["depreciation_end_year"])
        props.append(p)
    hc = {k: num(v) for k, v in kv(wb["healthcare"]).items()}
    data = dict(country=country, people=people, start=int(num(hh["start_year"])), kids=kids, household=household,
                kids_lines=kids_lines, other_lines=other_lines, props=props, hc=hc, tax=read_tax(wb["tax_tables"], country),
                joint_share=num(hh.get("joint_share_p1"), 0.5), tfsa_room0=num(hh.get("tfsa_room")))
    return data


def read_model(ws, data):
    """One Model tab: (scenario inputs, its year table {column key: [values]}, its results block)."""
    country = data["country"]
    lay = model_layout(country)
    for title, r in lay["sections"].items():
        if ws.cell(r, 1).value != title:
            raise SystemExit(f"'{ws.title}': expected the section '{title}' in row {r} - were rows inserted or deleted?")
    g = lambda key: ws.cell(lay["inputs"][key], 2).value  # noqa: E731
    S = {}
    for key in lay["inputs"]:
        v = g(key)
        S[key] = year(v) if key in ("p1_retire", "p2_retire", "ss_cut_year", "move_year", "conv_from", "conv_to") else num(v)
    B = BUCKETS[country]
    start_balances = dict(taxable=num(g("taxable_0")), basis=num(g("basis_0")), pre1=num(g("pre1_0")), pre2=num(g("pre2_0")),
                          free=num(g("free_0")), hsa=num(g("hsa_0")) if "hsa" in B else 0.0, college=num(g("college_0")))
    S["props"] = []
    factors = []
    for i in range(N_PROPS):
        r = lay["prop0"] + i
        d = dict(sell=year(ws.cell(r, 2).value), change=year(ws.cell(r, 3).value), rent=num(ws.cell(r, 4).value),
                 costs=num(ws.cell(r, 5).value))
        S["props"].append(d)
        p = data["props"][i]
        formula = p["value_now"] * (1 + p["appreciation"]) ** (d["sell"] - data["start"] + 1) if d["sell"] else 0.0
        typed = num(ws.cell(r, 6).value)
        factors.append(typed / formula if d["sell"] and formula else 1.0)
    S["sale_price_factor"] = factors
    S["oneoffs"] = []
    for i in range(N_EVENTS):
        r = lay["ev0"] + i
        if ws.cell(r, 2).value not in (None, "") and ws.cell(r, 4).value:
            S["oneoffs"].append(dict(year=year(ws.cell(r, 2).value), amount=num(ws.cell(r, 3).value), type=ws.cell(r, 4).value))
    letters = column_letters(country)
    table = {k: [] for k in letters}
    r = lay["first"]
    while ws[f"{letters['year']}{r}"].value not in (None, ""):
        for k, L in letters.items():
            table[k].append(ws[f"{L}{r}"].value)
        r += 1
    if not table["year"]:
        raise SystemExit(f"'{ws.title}' has no calculated values - open and save it in Excel/LibreOffice, or install LibreOffice")
    years = [int(y) for y in table["year"]]
    S["returns"] = {y: num(v) for y, v in zip(years, table["ret_override"]) if v not in (None, "")}
    S["spend_mult"] = {y: num(v) for y, v in zip(years, table["spend_mult"]) if v not in (None, "")}
    results = {key: ws.cell(RESULTS_ROW0 + i, 7).value for i, (key, _, _, _) in enumerate(RESULTS)}
    return S, dict(data, start_balances=start_balances, end=years[-1]), table, results


def read_plan(path):
    """{'data': shared inputs, 'models': {tab: (scenario, data_for_tab, table, results)}, 'wb': the workbook, ...}"""
    calc, fresh = recalculated(path)
    wb = openpyxl.load_workbook(calc, data_only=True)
    data = read_shared(wb)
    models = {ws.title: read_model(ws, data) for ws in wb.worksheets if ws.title.startswith("Model - ")}
    return dict(data=data, models=models, wb=wb, recalculated=fresh)


def headline(res, data):
    """The Model tab's results block, recomputed from a model.run() result."""
    by = {r["year"]: r for r in res}
    born = data["people"][0]["born"]
    short = next((r["year"] for r in res if r["shortfall"] > 1), None)
    get = lambda age, k: by[born + age][k] if born + age in by else ""  # noqa: E731
    return dict(nw70=get(70, "net_worth_today"), nw80=get(80, "net_worth_today"), nw85=get(85, "net_worth_today"),
                nw90=get(90, "net_worth_today"), nw95=get(95, "net_worth_today"),
                inv90=max(0.0, get(90, "investments_today")) if born + 90 in by else "",
                short=short if short else "never", tax=sum(r["income_tax"] / r["infl_f"] for r in res))


def compare_headline(sheet, ours):
    bad = []
    for k, want in sheet.items():
        if k not in ours or k in ("living1", "rate1"):
            continue
        got = ours[k]
        if isinstance(want, (int, float)) and isinstance(got, (int, float)):
            if abs(want - got) > max(1.0, 1e-4 * abs(want)):
                bad.append(f"{k}: sheet {want:,.0f} vs model {got:,.0f}")
        elif str(want if want is not None else "") != str(got):
            bad.append(f"{k}: sheet {want} vs model {got}")
    return bad


# ---------------------------------------------------------------------------------------------- editing

def input_cells(country):
    """{input key: cell} on a Model tab, including the property decisions (prop1_sell ...) and one-off events."""
    lay = model_layout(country)
    cells = {k: f"B{r}" for k, r in lay["inputs"].items()}
    for i in range(N_PROPS):
        for col, k in zip("BCDEF", ("sell", "change", "rent", "costs", "sale_price")):
            cells[f"prop{i + 1}_{k}"] = f"{col}{lay['prop0'] + i}"
    for i in range(N_EVENTS):
        for col, k in zip("ABCD", ("desc", "year", "amount", "type")):
            cells[f"event{i + 1}_{k}"] = f"{col}{lay['ev0'] + i}"
    return cells


def parse_value(v):
    """A value typed on the command line: '' = blank, '=...' = formula, '5%' = 0.05, numbers, else text."""
    v = v.strip()
    if v == "":
        return None
    if v.startswith("="):
        return v
    if v.endswith("%"):
        return float(v[:-1]) / 100
    for conv in (int, float):
        try:
            return conv(v)
        except ValueError:
            pass
    return v


def set_inputs(path, tab, changes):
    """Change inputs on one Model tab (keys from input_cells). Backs up first; returns [(key, cell, old, new)]."""
    from update_workbook import backup
    wb = openpyxl.load_workbook(path)
    country = str(kv(wb["household"])["country"]).strip().upper()
    if tab not in wb.sheetnames:
        raise SystemExit(f"no tab named {tab!r}")
    cells = input_cells(country)
    bad = [k for k in changes if k not in cells]
    if bad:
        raise SystemExit(f"unknown input(s): {', '.join(bad)} - `./plan inputs <folder>` lists them")
    dest = backup(path)
    ws, done = wb[tab], []
    for k, v in changes.items():
        old = ws[cells[k]].value
        ws[cells[k]] = v
        done.append((k, cells[k], old, v))
    wb.save(path)
    return done, dest


def duplicate_tab(path, src, name):
    from update_workbook import backup
    if not name.startswith("Model - "):
        name = "Model - " + name
    wb = openpyxl.load_workbook(path)
    if name in wb.sheetnames:
        raise SystemExit(f"{name!r} already exists")
    dest = backup(path)
    ws = wb.copy_worksheet(wb[src])
    ws.title = name
    ws["A1"] = "Retirement model - " + name.replace("Model - ", "")
    for dv in wb[src].data_validations.dataValidation:     # dropdowns aren't copied by openpyxl
        ws.add_data_validation(copy.copy(dv))
    wb.move_sheet(ws, offset=-(len(wb.sheetnames) - 1 - wb.sheetnames.index(src) - 1))
    wb.save(path)
    return name, dest
