#!/usr/bin/env python3
"""Create a new plan workbook (.xlsx) from a folder of input CSVs.

Every model number is a live formula: edit a yellow cell and everything downstream recalculates, in Excel,
LibreOffice or Google Sheets (File > Import). The "Model - ..." tab is a scenario; duplicate it to try another.
Scenario inputs live on the Model tab; shared facts (household, budgets, properties, healthcare, accounts, tax law)
live on their own tabs and feed every scenario.

This is for starting a plan. Once it exists, the workbook is the plan: edit it (or have Claude edit it), and use
update_workbook.py to bring in new statements - don't rebuild over it.

Usage: build_workbook.py <inputs_folder> <output.xlsx>
       (inputs_folder: like examples/us or examples/ca)
"""
import csv
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("build_workbook.py needs openpyxl: pip install openpyxl")

import tax_tables as TX
from model import ONEOFF_TYPES

DATA = Path(__file__).resolve().parent.parent / "data"
INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="DDEBF7")
BAND = PatternFill("solid", fgColor="1F4E78")
SUBTLE = PatternFill("solid", fgColor="F2F2F2")
KEY = PatternFill("solid", fgColor="E2EFDA")
BOLD = Font(bold=True)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=14)
GREY = Font(color="808080", italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")
MONEY, PCT, NUM3, YEAR = "#,##0", "0.0%", "0.000", "0"

MODEL_TAB = "Model - Example"
N_PROPS, N_EVENTS, N_KIDS = 5, 6, 8
BUCKETS = {  # model account -> the bucket names used on the assets tab
    "US": dict(taxable="taxable", pre1="pretax_p1", pre2="pretax_p2", free="roth", hsa="hsa", college="college_529"),
    "CA": dict(taxable="non_registered", pre1="rrsp_p1", pre2="rrsp_p2", free="tfsa", college="resp"),
}
BUCKET_HELP = {
    "US": [("taxable", "Taxable brokerage accounts, bank and money-market cash"),
           ("pretax_p1", "Person 1's 401(k), 403(b), traditional IRA, other pre-tax accounts"),
           ("pretax_p2", "Person 2's pre-tax accounts"),
           ("roth", "Roth IRAs and Roth 401(k)s (either person)"),
           ("hsa", "Health savings accounts"),
           ("college_529", "529 plans and other money set aside for the kids' college (not counted as yours)"),
           ("not_modeled", "Anything the model should ignore (cars, collectibles, a business)")],
    "CA": [("non_registered", "Non-registered investment accounts, bank and savings cash"),
           ("rrsp_p1", "Person 1's RRSP, RRIF, LIRA, defined-contribution pension"),
           ("rrsp_p2", "Person 2's RRSP/RRIF/LIRA"),
           ("tfsa", "TFSAs (either person)"),
           ("resp", "RESPs (the kids' education money; not counted as yours)"),
           ("not_modeled", "Anything the model should ignore (cars, collectibles, a business)")],
}
PROP_COLS = ["property", "type", "value_now", "appreciation", "selling_cost_pct", "cost_basis", "carrying_costs",
             "property_tax", "mortgage_balance", "mortgage_rate", "mortgage_payment", "rent", "cash_costs",
             "depreciation", "depreciation_end_year", "accum_depreciation", "notes"]


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def read_kv(path):
    return {r["key"]: r["value"] for r in csv.DictReader(open(path, newline="", encoding="utf-8-sig"))}


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    v = v.strip()
    if v == "":
        return None
    try:
        return float(v) if any(ch in v for ch in ".eE") else int(v)
    except ValueError:
        return v


def write_table(ws, rows, input_cols=(), money_cols=(), pct_cols=(), widths=None, start_row=1, extra_rows=0):
    """A CSV-like table; header bold/blue, input columns yellow (also `extra_rows` blank rows below)."""
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            v = num(val) if r_i else val
            cell = ws.cell(start_row + r_i, c_i + 1, v)
            if r_i == 0:
                cell.font, cell.fill, cell.alignment = BOLD, HEAD, WRAP
            else:
                if c_i in money_cols:
                    cell.number_format = MONEY
                if c_i in pct_cols:
                    cell.number_format = PCT
    for r_i in range(1, len(rows) + extra_rows):
        for c_i in input_cols:
            ws.cell(start_row + r_i, c_i + 1).fill = INPUT
    for i, w in enumerate(widths or []):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)


# ================================================================================ Model tab layout (shared with workbook.py)

def input_spec(country):
    """Model tab inputs, top to bottom: ("section", title) or (key, label, default, fmt, note, kind).
    kind: "in" = yellow input, "calc" = formula filled in by the builder. Labels may contain {P1}/{P2}."""
    us = country == "US"
    s = [("section", "WORK")]
    s += [("p1_retire", "{P1}'s first year without a paycheck", None, YEAR, "The year after the last working year", "in"),
          ("p2_retire", "{P2}'s first year without a paycheck", None, YEAR, "", "in"),
          ("p1_pay", "{P1}'s pay per year while working (today's $)", 0, MONEY, "Salary + bonus + stock vesting", "in"),
          ("p2_pay", "{P2}'s pay per year while working", 0, MONEY, "", "in"),
          ("p1_contrib", "{P1}'s " + ("401(k)/403(b) contributions (pre-tax)" if us else "RRSP / pension contributions"), 0, MONEY, "While working", "in"),
          ("p2_contrib", "{P2}'s " + ("401(k)/403(b) contributions (pre-tax)" if us else "RRSP / pension contributions"), 0, MONEY, "", "in"),
          ("p1_match", "Employer match / pension contributions for {P1}", 0, MONEY, "", "in"),
          ("p2_match", "Employer match / pension contributions for {P2}", 0, MONEY, "", "in")]
    if us:
        s += [("hsa_contrib", "HSA contributions (while either works)", 0, MONEY, "Family limit + catch-up", "in")]
    s += [("withholding_rate", "Income tax withheld from pay", 0.25, PCT, "Share of taxable pay; the difference is settled next April", "in"),
          ("tax_due_first", "Tax owed for last year, paid in April of the first year", 0, MONEY, "Negative = a refund", "in"),
          ("section", "ECONOMY"),
          ("inflation", "Inflation", 0.03, PCT, "Living costs, pay, pensions, tax brackets", "in"),
          ("return_default", "Investment return (per year, before inflation)", 0.05, PCT, "Override any year in the yellow 'Return override' column below", "in"),
          ("dividend_yield", "Dividend yield on " + ("taxable" if us else "non-registered") + " accounts", 0.015, PCT, "Paid out as cash each year", "in")]
    if not us:
        s += [("eligible_div_share", "Share of those dividends that are eligible Canadian dividends", 0.5, PCT, "The rest (foreign, interest) is taxed as ordinary income", "in")]
    s += [("med_infl", "Healthcare cost growth", "=HC:medical_inflation", PCT, "From the healthcare tab", "in"),
          ("col_infl", "College cost growth", 0.04, PCT, "", "in"),
          ("ltc_infl", "Long-term care cost growth", "=HC:ltc_inflation", PCT, "From the healthcare tab", "in")]
    if us:
        s += [("section", "SOCIAL SECURITY"),
              ("p1_claim_age", "{P1} starts Social Security at age", 67, "0", "62-70", "in"),
              ("p2_claim_age", "{P2} starts Social Security at age", 67, "0", "62-70. A spouse can get up to half the other's benefit once both claim", "in"),
              ("ss_early_cut", "Benefit reduction for stopping work early", 0.0, PCT, "SSA statements assume you work until claiming; zero-earning years lower the average (often 1-5%)", "in"),
              ("ss_cut_pct", "Across-the-board benefit cut", 0.19, PCT, "If Congress doesn't act, the trust fund pays ~81% from about 2034", "in"),
              ("ss_cut_year", "Cut starts in", 2034, YEAR, "Set the cut to 0% to assume Congress fixes it", "in"),
              ("p1_claim_year", "{P1}'s claim year", "calc", YEAR, "", "calc"),
              ("p2_claim_year", "{P2}'s claim year", "calc", YEAR, "", "calc"),
              ("p1_ss_factor", "{P1}'s benefit vs full retirement age", "calc", NUM3, "Early claims are reduced, later ones get 8%/yr", "calc"),
              ("p2_ss_factor", "{P2}'s benefit vs full retirement age", "calc", NUM3, "", "calc"),
              ("p1_spousal_factor", "{P1}'s spousal benefit factor", "calc", NUM3, "Spousal benefits don't grow after full retirement age", "calc"),
              ("p2_spousal_factor", "{P2}'s spousal benefit factor", "calc", NUM3, "", "calc"),
              ("section", "WHERE YOU LIVE (STATE TAX)"),
              ("move_year", "Year you move to another state", None, YEAR, "Blank = stay. The state brackets are on the tax_tables tab", "in"),
              ("other_state_rate", "The new state's income tax rate", 0.0, PCT, "Flat rate on income except Social Security; 0% for FL, TX, TN, NV, WA...", "in")]
    else:
        s += [("section", "CPP AND OAS"),
              ("p1_cpp_age", "{P1} starts CPP at age", 65, "0", "60-70: -0.6%/month before 65, +0.7%/month after", "in"),
              ("p2_cpp_age", "{P2} starts CPP at age", 65, "0", "", "in"),
              ("p1_oas_age", "{P1} starts OAS at age", 65, "0", "65-70: +0.6%/month deferred", "in"),
              ("p2_oas_age", "{P2} starts OAS at age", 65, "0", "", "in"),
              ("p1_cpp_factor", "{P1}'s CPP vs age 65", "calc", NUM3, "", "calc"),
              ("p2_cpp_factor", "{P2}'s CPP vs age 65", "calc", NUM3, "", "calc"),
              ("p1_oas_factor", "{P1}'s OAS deferral increase", "calc", NUM3, "", "calc"),
              ("p2_oas_factor", "{P2}'s OAS deferral increase", "calc", NUM3, "", "calc")]
    s += [("section", "WORKPLACE PENSIONS"),
          ("p1_pension", "{P1}'s defined-benefit pension per year (today's $)", 0, MONEY, "0 = none. Your pension statement's estimate", "in"),
          ("p1_pension_age", "... starts at age", 65, "0", "", "in"),
          ("p2_pension", "{P2}'s defined-benefit pension per year (today's $)", 0, MONEY, "", "in"),
          ("p2_pension_age", "... starts at age", 65, "0", "", "in"),
          ("pension_indexing", "Pension increases, as a share of inflation", 0.0, "0%", "0% = fixed dollar amount, 100% = fully indexed", "in"),
          ("pension_survivor", "Share that continues to the survivor", 0.6, "0%", "Joint-and-survivor option; used in the Monte Carlo's lifetime runs", "in"),
          ("section", "HEALTHCARE AND LONG-TERM CARE"),
          ("p1_ltc_age", "{P1} needs long-term care from age", 85, "0", "Paid from savings" + (" (HSA first)" if us else ""), "in"),
          ("p1_ltc_years", "... for this many years", 2, "0", "Averages: men ~2 years, women ~3.5", "in"),
          ("p2_ltc_age", "{P2} needs long-term care from age", 87, "0", "Overlapping years double the cost", "in"),
          ("p2_ltc_years", "... for this many years", 3, "0", "", "in"),
          ("ltc_cost", "Long-term care cost per person-year (today's $)", "=HC:ltc_annual_cost", MONEY, "From the healthcare tab", "in"),
          ("section", "TAX MOVES")]
    if us:
        s += [("conv_amount", "Roth conversion per year (today's $)", 0, MONEY, "Moves pre-tax money to the Roth; taxed as income that year", "in"),
              ("conv_from", "... from year", None, YEAR, "Blank = no conversions", "in"),
              ("conv_to", "... through year", None, YEAR, "", "in"),
              ("charity", "Charitable giving per year (today's $)", 0, MONEY, "Deductible if you itemize; the spending itself belongs in the budget", "in")]
    else:
        s += [("conv_amount", "Extra RRSP/RRIF withdrawal per year (today's $)", 0, MONEY, "An 'RRSP meltdown': taxed now, the cash fills the TFSA or pays the bills", "in"),
              ("conv_from", "... from year", None, YEAR, "Blank = none", "in"),
              ("conv_to", "... through year", None, YEAR, "", "in")]
    s += [("section", "STARTING POINT (1 January of the first year) - from the assets tab; type over to freeze")]
    if us:
        s += [("taxable_0", "Taxable investments + cash", "bucket:taxable", MONEY, "", "calc"),
              ("basis_0", "... their cost basis", "basis:taxable", MONEY, "", "calc"),
              ("pre1_0", "{P1}'s pre-tax accounts", "bucket:pre1", MONEY, "", "calc"),
              ("pre2_0", "{P2}'s pre-tax accounts", "bucket:pre2", MONEY, "", "calc"),
              ("free_0", "Roth accounts", "bucket:free", MONEY, "", "calc"),
              ("hsa_0", "HSAs", "bucket:hsa", MONEY, "", "calc"),
              ("college_0", "529s / college funds", "bucket:college", MONEY, "", "calc"),
              ("p1_rmd_age", "{P1}'s required withdrawals start at", "calc", "0", "73 if born 1951-59, 75 if born 1960 or later", "calc"),
              ("p2_rmd_age", "{P2}'s required withdrawals start at", "calc", "0", "", "calc")]
    else:
        s += [("taxable_0", "Non-registered investments + cash", "bucket:taxable", MONEY, "", "calc"),
              ("basis_0", "... their adjusted cost base", "basis:taxable", MONEY, "", "calc"),
              ("pre1_0", "{P1}'s RRSP/RRIF", "bucket:pre1", MONEY, "", "calc"),
              ("pre2_0", "{P2}'s RRSP/RRIF", "bucket:pre2", MONEY, "", "calc"),
              ("free_0", "TFSAs", "bucket:free", MONEY, "", "calc"),
              ("college_0", "RESPs", "bucket:college", MONEY, "", "calc")]
    return s


def model_layout(country):
    """Row numbers for everything on a Model tab (the builder writes them; workbook.py reads them)."""
    rows, r = {}, 4
    sections = {}
    for item in input_spec(country):
        if item[0] == "section":
            if sections:
                r += 1
            sections[item[1]] = r
        else:
            rows[item[0]] = r
        r += 1
    r += 1
    prop_head = r
    prop0 = r + 1
    prop_note = prop0 + N_PROPS
    ev_head = prop_note + 2
    ev0 = ev_head + 1
    band = ev0 + N_EVENTS + 1
    return dict(inputs=rows, sections=sections, prop_head=prop_head, prop0=prop0, prop_note=prop_note,
                ev_head=ev_head, ev0=ev0, band=band, head=band + 1, first=band + 2)


RESULTS = [("nw70", "Net worth when {P1} is 70", "nw", 70), ("nw80", "... at 80", "nw", 80),
           ("nw85", "... at 85", "nw", 85), ("nw90", "... at 90", "nw", 90), ("nw95", "... at 95", "nw", 95),
           ("inv90", "Investments when {P1} is 90", "inv", 90), ("short", "First year investments can't cover costs", "short", None),
           ("living1", "First year's living costs", "living1", None), ("rate1", "First year's withdrawal rate", "rate1", None),
           ("tax", "Lifetime income tax (today's $)", "tax", None)]
RESULTS_ROW0 = 5   # results in column G from this row


# ================================================================================ the workbook

class Book:
    def __init__(self, inputs):
        self.L = Path(inputs)
        self.hh = read_kv(self.L / "household.csv")
        self.country = self.hh["country"].strip().upper()
        assert self.country in ("US", "CA"), "household.csv: country must be US or CA"
        self.us = self.country == "US"
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.H, self.HC, self.TT, self.BR = {}, {}, {}, {}

    # ------------------------------------------------------------------ shared tabs
    def household(self):
        ws = self.wb.create_sheet("household")
        ws["A1"], ws["B1"], ws["C1"] = "key", "value", "notes"
        for c in "ABC":
            ws[f"{c}1"].font, ws[f"{c}1"].fill = BOLD, HEAD
        keys = [("country", "US = United States, CA = Canada. Fixed when the workbook is built"),
                ("region", "State or province (its tax brackets are on the tax_tables tab)"),
                ("start_year", "First year the model covers. Balances on the assets tab are as of 1 January of it; "
                               "'today's dollars' are the year before"),
                ("p1_name", "Person 1"), ("p1_born", "Birth year"),
                ("p1_sex", "M or F: which column of the life table the Monte Carlo uses for person 1"),
                ("p2_name", "Person 2"), ("p2_born", "Birth year"), ("p2_sex", "M or F for person 2")]
        if self.us:
            keys += [("p1_ss_benefit", "Person 1's Social Security at full retirement age (67), per month, today's $ - from ssa.gov/myaccount"),
                     ("p2_ss_benefit", "Person 2's, on their own record")]
        else:
            keys += [("p1_cpp_at_65", "Person 1's CPP at 65, per month, today's $ - from My Service Canada Account"),
                     ("p2_cpp_at_65", "Person 2's CPP at 65"),
                     ("p1_oas_years", "Person 1's years in Canada after 18 (40 = full OAS)"),
                     ("p2_oas_years", "Person 2's years in Canada after 18"),
                     ("joint_share_p1", "Share of joint investment and rental income taxed to person 1 (usually 50%)"),
                     ("tfsa_room", "Unused TFSA room for the two of you on 1 January of the first year (incl. that year's new room)")]
        for i, (k, note) in enumerate(keys, start=2):
            ws.cell(i, 1, k)
            c = ws.cell(i, 2, num(self.hh.get(k, "")))
            c.fill = INPUT
            if k == "joint_share_p1":
                c.number_format = "0%"
            if k.endswith("_benefit") or k.endswith("_at_65") or k == "tfsa_room":
                c.number_format = MONEY
            ws.cell(i, 3, note)
            self.H[k] = f"household!$B${i}"
        ws.cell(2, 2).fill = SUBTLE
        r0 = len(keys) + 4
        ws.cell(r0 - 1, 1, "Kids").font = BOLD
        kids = read_csv(self.L / "kids.csv")
        head = ["name", "born", "on_health_plan_through" if self.us else "(not used)", "notes"]
        for j, h in enumerate(head):
            c = ws.cell(r0, j + 1, h)
            c.font, c.fill = BOLD, HEAD
        for i in range(N_KIDS):
            row = kids[i + 1] if i + 1 < len(kids) else []
            for j in range(4):
                c = ws.cell(r0 + 1 + i, j + 1, num(row[j]) if j < len(row) else None)
                if j < 3:
                    c.fill = INPUT
        ws.cell(r0 + N_KIDS + 1, 1, ("Kids' costs are on the kids_budget tab (one line per kid and cost)."
                                     + (" Health-plan years matter once you buy your own coverage." if self.us else "")))
        self.H["kids_through"] = f"household!$C${r0 + 1}:$C${r0 + N_KIDS}"
        self.kids_rows = (r0 + 1, r0 + N_KIDS)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 110
        self.P1, self.P2 = self.H["p1_name"], self.H["p2_name"]

    def budget(self):
        ws = self.wb.create_sheet("budget")
        rows = read_csv(self.L / "budget.csv")
        write_table(ws, rows, input_cols=(0, 5), money_cols=(4, 5), widths=[12, 16, 22, 30, 13, 15, 22, 8, 70])
        ws["K1"], ws["K2"] = "Household living budget used by the model", "All lines, 12-month actual"
        ws["K1"].font = BOLD
        ws["L1"] = '=SUMIF($A$2:$A$300,"household",$F$2:$F$300)'
        ws["L2"] = "=SUM($E$2:$E$300)"
        for ref in ("L1", "L2"):
            ws[ref].number_format = MONEY
        ws["L1"].font = BOLD
        ws["K4"] = ("model_line: 'household' lines are summed and grow with inflation. Other lines (housing, healthcare, "
                    "taxes, kids) are modeled on their own tabs and shown here only for reference.")
        ws["K5"] = "Edit column F (retirement_budget) to test a different budget; column E is what you actually spent."
        ws["K4"].alignment = ws["K5"].alignment = WRAP
        ws.column_dimensions["K"].width = 44
        ws.column_dimensions["L"].width = 14
        self.budget_rng = ("budget!$A$2:$A$300", "budget!$F$2:$F$300")

    def kids_budget(self):
        ws = self.wb.create_sheet("kids_budget")
        rows = read_csv(self.L / "kids_budget.csv")
        write_table(ws, rows, input_cols=(0, 1, 2, 3, 4, 5, 6), money_cols=(2,), widths=[10, 44, 14, 10, 10, 10, 14, 70],
                    extra_rows=max(0, 40 - len(rows)))
        dv1 = DataValidation(type="list", formula1='"general,college"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"parents,college_funds"', allow_blank=True)
        ws.add_data_validation(dv1)
        ws.add_data_validation(dv2)
        dv1.add("F2:F40")
        dv2.add("G2:G40")
        ws.cell(42, 1, ("One line per kid and cost, per year in today's dollars, for the years start_year-end_year. "
                        "inflation: 'college' lines rise with college cost growth, 'general' with inflation. paid_from: "
                        "'college_funds' lines are paid by the " + ("529s" if self.us else "RESPs") + " first and only the "
                        "shortfall comes from you; 'parents' lines come straight from you. Add rows above row 41."))
        for k, col in (("kid", "A"), ("amt", "C"), ("start", "D"), ("end", "E"), ("infl", "F"), ("paid", "G")):
            setattr(self, "kb_" + k, f"kids_budget!${col}$2:${col}$40")

    def other_budget(self):
        ws = self.wb.create_sheet("other_budget")
        rows = read_csv(self.L / "other_budget.csv")
        write_table(ws, rows, input_cols=(0, 1, 2, 3), money_cols=(1,), widths=[46, 14, 10, 10, 80],
                    extra_rows=max(0, 30 - len(rows)))
        ws.cell(32, 1, "Big irregular costs averaged per year (renovations, cars, gifts), in today's dollars, for the "
                       "years start_year-end_year. Add rows above row 31.")
        self.ob = {k: f"other_budget!${c}$2:${c}$30" for k, c in (("amt", "B"), ("start", "C"), ("end", "D"))}

    def properties(self):
        ws = self.wb.create_sheet("properties")
        rows = read_csv(self.L / "properties.csv")
        assert rows[0] == PROP_COLS, f"properties.csv columns must be: {','.join(PROP_COLS)}"
        rows += [[""] * len(PROP_COLS)] * (N_PROPS + 1 - len(rows))
        write_table(ws, rows, input_cols=tuple(range(0, 16)), money_cols=(2, 5, 6, 7, 8, 10, 11, 12, 13, 15),
                    pct_cols=(3, 4, 9), widths=[30, 12] + [12] * 14 + [80])
        for r in range(2, 2 + N_PROPS):
            ws.cell(r, 10).number_format = "0.00%"
        ws.cell(2, 2).fill = SUBTLE
        notes = [
            "Row 2 is your home (principal residence); rows 3-6 are other properties (rentals, a cottage). Sell / rent "
            "decisions are made on each Model tab. Amounts are today's dollars.",
            "value_now: market value; appreciation: growth per year. carrying_costs: property tax + condo/HOA/strata fees "
            "+ insurance + utilities you want counted as housing (for the home; if you rent, put the rent here and leave "
            "value_now blank). property_tax: the part that's property tax" + (" (deductible up to the SALT cap)." if self.us else " (not used in Canada)."),
            "mortgage_*: the home's mortgage on 1 January of the first year (balance, interest rate, payments per year). "
            "Other properties are assumed mortgage-free: include any loan payments in cash_costs.",
            "Other properties: rent and cash_costs per year; depreciation" + (" " if self.us else " (capital cost allowance) ")
            + "claimed per year until depreciation_end_year, and accum_depreciation already claimed. cost_basis is what you paid plus improvements.",
        ]
        for i, t in enumerate(notes):
            ws.cell(9 + i, 1, t)
        self.pc = {k: get_column_letter(i + 1) for i, k in enumerate(PROP_COLS)}

    def healthcare(self):
        ws = self.wb.create_sheet("healthcare")
        rows = read_csv(self.L / "healthcare.csv")
        write_table(ws, rows, input_cols=(1,), widths=[30, 12, 16, 80, 50])
        for r_i, row in enumerate(rows[1:], start=2):
            self.HC[row[0]] = f"healthcare!$B${r_i}"
            if "inflation" in row[0]:
                ws.cell(r_i, 2).number_format = PCT

    def assets(self):
        ws = self.wb.create_sheet("assets")
        rows = read_csv(self.L / "assets.csv")
        n = 150
        write_table(ws, rows, input_cols=(0, 1, 2, 3, 4), money_cols=(2, 4), widths=[36, 16, 14, 11, 14, 70],
                    extra_rows=n - len(rows))
        buckets = [b for b, _ in BUCKET_HELP[self.country]]
        dv = DataValidation(type="list", formula1='"' + ",".join(buckets) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"B2:B{n}")
        ws["H1"], ws["I1"], ws["J1"], ws["K1"] = "bucket", "total", "cost basis", "what goes in it"
        for c in "HIJK":
            ws[f"{c}1"].font, ws[f"{c}1"].fill = BOLD, HEAD
        for i, (b, help_) in enumerate(BUCKET_HELP[self.country], start=2):
            ws.cell(i, 8, b)
            ws.cell(i, 9, f'=SUMIF($B$2:$B${n},"{b}",$C$2:$C${n})').number_format = MONEY
            ws.cell(i, 11, help_)
        ws.cell(2, 10, f'=SUMIF($B$2:$B${n},"{BUCKETS[self.country]["taxable"]}",$E$2:$E${n})').number_format = MONEY
        ws.cell(len(BUCKET_HELP[self.country]) + 3, 8, ("One row per account, balance as of 1 January of the first year. "
                                                        "cost_basis matters only for taxable accounts (blank = same as value, "
                                                        "e.g. cash)."))
        ws.column_dimensions["H"].width = 16
        ws.column_dimensions["I"].width = 14
        ws.column_dimensions["J"].width = 14
        ws.column_dimensions["K"].width = 70
        self.as_rng = dict(bucket=f"assets!$B$2:$B${n}", value=f"assets!$C$2:$C${n}", basis=f"assets!$E$2:$E${n}")
        self.as_basis_total = "assets!$J$2"

    def tax_tables(self):
        ws = self.wb.create_sheet("tax_tables")
        us = self.us
        ws["A1"] = ("Tax law used by the model (married filing jointly). Yellow cells can be edited; indexed amounts grow "
                    "with inflation after TAX_YEAR." if us else
                    "Tax law used by the model (each spouse files separately; pension income is split). Yellow cells "
                    "can be edited; indexed amounts grow with inflation after TAX_YEAR.")
        ws["A1"].font = BOLD
        ws["A2"] = ("Sources: IRS Rev. Proc. 2025-32 (2026 amounts, One Big Beautiful Bill Act); SSA and CMS 2026 figures; "
                    "state brackets from the state's tax department. Update each January." if us else
                    "Sources: CRA 2026 indexation (canada.ca, tax rates and payroll deduction tables); BC Budget 2026 "
                    "(lowest rate 5.6%, indexing paused 2027-2030); ESDC CPP/OAS amounts (Jan 2026). Update each January.")
        scalars = TX.US_SCALARS if us else TX.CA_SCALARS
        region = self.hh.get("region", "")
        for c, h in zip("ABC", ("name", "value", "notes")):
            ws[f"{c}3"], ws[f"{c}3"].font, ws[f"{c}3"].fill = h, BOLD, HEAD
        for i, (k, v, note) in enumerate(scalars, start=4):
            if k in ("STATE_NAME", "PROV_NAME") and region:
                v = region
            ws.cell(i, 1, k)
            c = ws.cell(i, 2, v)
            c.fill = INPUT
            if isinstance(v, float) and v < 1:
                c.number_format = "0.00%" if v < 0.1 else "0.0%"
            ws.cell(i, 3, note)
            self.TT[k] = f"tax_tables!$B${i}"
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 70

        def block(r0, c0, title, header, data, nrows=None, input_cols=()):
            ws.cell(r0, c0, title).font = BOLD
            for j, h in enumerate(header):
                c = ws.cell(r0 + 1, c0 + j, h)
                c.font, c.fill = BOLD, HEAD
            for i in range(nrows or len(data)):
                for j in range(len(header)):
                    v = data[i][j] if i < len(data) else None
                    c = ws.cell(r0 + 2 + i, c0 + j, v)
                    if j in input_cols:
                        c.fill = INPUT
            return r0 + 2, r0 + 1 + (nrows or len(data))

        col = 5
        brackets = TX.US_BRACKETS if us else TX.CA_BRACKETS
        for key, (title, indexed, br) in brackets.items():
            if key in ("state", "prov"):
                title = f"{title}: {region}" if region else title
            r0 = 3
            a, b = block(r0, col, title, ["over", "rate", "step"], [(lo, rate, None) for lo, rate in br],
                         nrows=TX.BRACKET_ROWS, input_cols=(0, 1))
            L = [get_column_letter(col + k) for k in range(3)]
            for rr in range(a, b + 1):
                prev = f"{L[1]}{rr - 1}" if rr > a else "0"
                ws[f"{L[2]}{rr}"] = f'=IF({L[1]}{rr}="",0,{L[1]}{rr}-{prev})' if rr == a else \
                    f'=IF({L[1]}{rr}="",0,{L[1]}{rr}-{L[1]}{rr - 1})'
                ws[f"{L[1]}{rr}"].number_format = "0.00%"
                ws[f"{L[2]}{rr}"].number_format = "0.00%"
                ws[f"{L[0]}{rr}"].number_format = MONEY
            self.BR[key] = tuple(f"tax_tables!${L[k]}${a}:${L[k]}${b}" for k in range(3))
            self.BR[key + "_pos"] = (col, a, b)
            ws.column_dimensions[L[0]].width = 12
            ws.column_dimensions[L[1]].width = 8
            ws.column_dimensions[L[2]].width = 8
            ws.column_dimensions[get_column_letter(col + 3)].width = 2
            col += 4
        if us:
            data, prev = [], 0.0
            for thr, b, d in TX.US_IRMAA:
                data.append((thr, b, d, None))
            a, b = block(3, col, "Medicare IRMAA (per person per month)", ["MAGI over", "Part B add-on", "Part D add-on", "step"], data,
                         input_cols=(0, 1, 2))
            L = [get_column_letter(col + k) for k in range(4)]
            for rr in range(a, b + 1):
                ws[f"{L[3]}{rr}"] = f"={L[1]}{rr}+{L[2]}{rr}" + (f"-{L[1]}{rr - 1}-{L[2]}{rr - 1}" if rr > a else "")
            self.BR["irmaa"] = (f"tax_tables!${L[0]}${a}:${L[0]}${b}", f"tax_tables!${L[3]}${a}:${L[3]}${b}")
            self.BR["irmaa_pos"] = (col, a, b)
            col += 5
            a, b = block(3, col, "SALT cap by year", ["year", "cap"], TX.US_SALT, input_cols=(0, 1))
            L = [get_column_letter(col + k) for k in range(2)]
            self.BR["salt"] = (f"tax_tables!${L[0]}${a}:${L[0]}${b}", f"tax_tables!${L[1]}${a}:${L[1]}${b}")
            self.BR["salt_pos"] = (col, a, b)
            col += 3
            a, b = block(3, col, "Uniform Lifetime Table (RMDs)", ["age", "factor"], sorted(TX.UNIFORM_LIFETIME.items()))
            L = [get_column_letter(col + k) for k in range(2)]
            self.BR["factors"] = (f"tax_tables!${L[0]}${a}:${L[0]}${b}", f"tax_tables!${L[1]}${a}:${L[1]}${b}")
            self.BR["factors_pos"] = (col, a, b)
        else:
            a, b = block(3, col, "RRIF minimum withdrawal (age on 1 January)", ["age", "minimum"], sorted(TX.RRIF_FACTORS.items()))
            L = [get_column_letter(col + k) for k in range(2)]
            for rr in range(a, b + 1):
                ws[f"{L[1]}{rr}"].number_format = "0.00%"
            self.BR["factors"] = (f"tax_tables!${L[0]}${a}:${L[0]}${b}", f"tax_tables!${L[1]}${a}:${L[1]}${b}")
            self.BR["factors_pos"] = (col, a, b)

    def data_tabs(self):
        ws = self.wb.create_sheet("historical_returns")
        rows = read_csv(DATA / "historical_returns.csv")
        ws.append(rows[0])
        for r in rows[1:]:
            ws.append([int(r[0])] + [float(v) for v in r[1:]])
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=5):
            for c in row:
                c.number_format = "0.0%"
        for c in ws[1]:
            c.font, c.fill = BOLD, HEAD
        ws["G1"] = ("US history: S&P 500 total return, 10-year Treasury, 3-month T-bill, CPI inflation. Source: Aswath "
                    "Damodaran, NYU Stern (histretSP.xls, January 2026). Add a row each January, or replace the columns "
                    "with another country's series (keep the headers).")
        ws = self.wb.create_sheet("life_table")
        rows = read_csv(DATA / "life_table_ssa_2023.csv")
        ws.append(rows[0])
        for r in rows[1:]:
            ws.append([int(r[0]), float(r[1]), float(r[2])])
        for c in ws[1]:
            c.font, c.fill = BOLD, HEAD
        ws["E1"] = ("Probability of dying within a year at each age. US Social Security period life table for 2023 (as "
                    "used in the 2026 Trustees Report). The Monte Carlo shifts it by the extra years of life expectancy "
                    "set on the monte_carlo tab.")

    def monte_carlo_tab(self):
        from monte_carlo import SETTINGS, settings_labels
        ws = self.wb.create_sheet("monte_carlo")
        ws["A1"] = "Monte Carlo: every scenario replayed through history"
        ws["A1"].font = TITLE
        ws["A2"] = ("Each run replaces the Model tab's flat return and inflation with a stretch of real history "
                    "(historical_returns tab) and draws lifetimes from the life_table tab. Change the yellow cells, then "
                    "run the Monte Carlo (see README). Results are written below, in today's dollars.")
        ws["A4"] = "SETTINGS"
        ws["A4"].font = BOLD
        names = {"P1": str(self.hh.get("p1_name", "Person 1")), "P2": str(self.hh.get("p2_name", "Person 2"))}
        for (row, _, default, note, fmt), label in zip(SETTINGS, settings_labels(names)):
            ws.cell(row, 1, label)
            c = ws.cell(row, 2, default)
            c.fill = INPUT
            if fmt:
                c.number_format = fmt
            ws.cell(row, 3, note)
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 14

    # ------------------------------------------------------------------ Model tab
    def model(self, title=MODEL_TAB):
        ws = self.wb.create_sheet(title, 0)
        self.ws = ws
        lay = model_layout(self.country)
        self.lay = lay
        scen = read_kv(self.L / "scenario.csv")
        ws["A1"] = "Retirement model - " + title.replace("Model - ", "")
        ws["A1"].font = TITLE
        ws["A2"] = ("Yellow cells are this scenario's inputs. To try another scenario: right-click this tab > Duplicate, "
                    "rename it 'Model - <name>', change its yellow cells. Shared facts (budgets, properties, accounts, "
                    "healthcare, tax law) are on the other tabs. Don't insert or delete rows on this tab.")
        ws["A2"].alignment = WRAP
        ws.merge_cells("A2:L2")
        ws.row_dimensions[2].height = 30
        I = {}
        self.I = I
        label = self.label
        B = BUCKETS[self.country]
        for item in input_spec(self.country):
            if item[0] == "section":
                r = lay["sections"][item[1]]
                c = ws.cell(r, 1, item[1])
                c.font, c.fill = WHITE_BOLD, BAND
                for cc in (2, 3):
                    ws.cell(r, cc).fill = BAND
                continue
            key, text, default, fmt, note, kind = item
            r = lay["inputs"][key]
            I[key] = f"$B${r}"
            ws.cell(r, 1, label(text))
            ws.cell(r, 3, note)
            if kind == "in":
                v = num(scen[key]) if key in scen else default
                if isinstance(v, str) and v.startswith("=HC:"):
                    v = "=" + self.HC[v[4:]]
                c = ws.cell(r, 2, v)
                c.fill = INPUT
            else:
                c = ws.cell(r, 2)
            if fmt:
                c.number_format = fmt
        # calculated inputs (need I filled in first)
        for item in input_spec(self.country):
            if item[0] == "section" or item[5] != "calc":
                continue
            key, default = item[0], item[2]
            c = ws[I[key].replace("$", "")]
            if isinstance(default, str) and default.startswith("bucket:"):
                c.value = f'=SUMIF({self.as_rng["bucket"]},"{B[default[7:]]}",{self.as_rng["value"]})'
            elif isinstance(default, str) and default.startswith("basis:"):
                c.value = f"={self.as_basis_total}"
            elif key.endswith("_claim_year"):
                p = key[:2]
                c.value = f"={self.H[p + '_born']}+{I[p + '_claim_age']}"
            elif key.endswith("_ss_factor"):
                p = key[:2]
                m = f"ROUND(({I[p + '_claim_age']}-{self.TT['SS_FULL_RETIREMENT_AGE']})*12,0)"
                c.value = f"=IF({m}>=0,1+0.08/12*MIN({m},36),1-5/900*MIN(-{m},36)-5/1200*MAX(0,-{m}-36))"
            elif key.endswith("_spousal_factor"):
                p = key[:2]
                m = f"ROUND(({I[p + '_claim_age']}-{self.TT['SS_FULL_RETIREMENT_AGE']})*12,0)"
                c.value = f"=IF({m}>=0,1,1-25/3600*MIN(-{m},36)-5/1200*MAX(0,-{m}-36))"
            elif key.endswith("_cpp_factor"):
                p = key[:2]
                m = f"ROUND(({I[p + '_cpp_age']}-65)*12,0)"
                c.value = f"=IF({m}>=0,1+0.007*MIN({m},60),1+0.006*MAX({m},-60))"
            elif key.endswith("_oas_factor"):
                p = key[:2]
                c.value = f"=1+0.006*MIN(60,MAX(0,ROUND(({I[p + '_oas_age']}-65)*12,0)))"
            elif key.endswith("_rmd_age"):
                p = key[:2]
                b = self.H[p + "_born"]
                c.value = f"=IF({b}<1951,72,IF({b}<1960,73,75))"

        # ---- property decisions
        r = lay["prop_head"]
        hdr = ["Property", "Sell in year", "Change use from year", "New rent / yr (today's $)", "New costs / yr (today's $)",
               "Sale price", "Net proceeds", "Taxable gain", "Recapture" if self.us else "CCA recapture"]
        for j, h in enumerate(hdr):
            c = ws.cell(r, 1 + j, h)
            c.font, c.fill, c.alignment = BOLD, HEAD, WRAP
        ws.row_dimensions[r].height = 30
        p0 = lay["prop0"]
        self.prop_rows = (p0, p0 + N_PROPS - 1)
        for i in range(N_PROPS):
            rr = p0 + i
            ws.cell(rr, 1, f'=IF(properties!$A${2 + i}="","",properties!$A${2 + i})')
            for col in (2, 3, 4, 5):
                v = num(scen.get(f"prop{i + 1}_{['sell', 'change', 'rent', 'costs'][col - 2]}", ""))
                c = ws.cell(rr, col, v)
                c.fill = INPUT
            for col in (4, 5):
                ws.cell(rr, col).number_format = MONEY
            ws.cell(rr, 2).number_format = ws.cell(rr, 3).number_format = YEAR
        ws.cell(p0, 3).fill = SUBTLE
        ws.cell(p0, 4).fill = SUBTLE
        ws.cell(lay["prop_note"], 1, ("Home: 'New costs' = what housing costs per year after you sell it (rent, or a new "
                                      "place's carrying costs; if you buy one, its price goes in one-off Expenses below and its value isn't counted in net worth). Other properties: "
                                      "'Change use' switches to the new rent/costs from that year (e.g. move into the "
                                      "cottage: rent 0, costs = its upkeep). Sale price is a formula; type over it to use "
                                      "your own estimate."))
        ws.merge_cells(start_row=lay["prop_note"], start_column=1, end_row=lay["prop_note"], end_column=9)
        ws.cell(lay["prop_note"], 1).alignment = WRAP
        ws.row_dimensions[lay["prop_note"]].height = 44

        # ---- one-off events
        r = lay["ev_head"]
        for j, h in enumerate(["One-off money in or out", "Year", "Amount (today's $)", "Type"]):
            c = ws.cell(r, 1 + j, h)
            c.font, c.fill = BOLD, HEAD
        dv = DataValidation(type="list", formula1='"' + ",".join(ONEOFF_TYPES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        e0 = lay["ev0"]
        self.ev_rows = (e0, e0 + N_EVENTS - 1)
        for i in range(N_EVENTS):
            rr = e0 + i
            for col, k in ((1, "desc"), (2, "year"), (3, "amount"), (4, "type")):
                v = scen.get(f"event{i + 1}_{k}", "")
                c = ws.cell(rr, col, num(v) if k in ("year", "amount") else (v or None))
                c.fill = INPUT
            ws.cell(rr, 2).number_format = YEAR
            ws.cell(rr, 3).number_format = MONEY
            dv.add(f"D{rr}")
        ws.cell(e0, 6, ("Severance, an inheritance, a wedding, a new car, buying a smaller home. Wages get payroll tax; "
                        "'Other taxable income' is taxed as ordinary income; 'Tax-free money in' isn't taxed; "
                        "'Expense' is spent that year."))
        ws.cell(e0, 6).font = GREY

        # ---- year table
        cols = self.model_columns()
        self.letters = {col[0]: get_column_letter(1 + j) for j, col in enumerate(cols)}
        first = lay["first"]
        self.first = first
        n_years = self.n_years()
        self.last = last = first + n_years - 1
        self.property_formulas()
        groups = []
        for j, col in enumerate(cols):
            if not groups or groups[-1][0] != col[2]:
                groups.append([col[2], j + 1, j + 1])
            else:
                groups[-1][2] = j + 1
        for g, c0, c1 in groups:
            cell = ws.cell(lay["band"], c0, g)
            cell.font, cell.fill = WHITE_BOLD, BAND
            for c in range(c0, c1 + 1):
                ws.cell(lay["band"], c).fill = BAND
            if c1 > c0:
                ws.merge_cells(start_row=lay["band"], start_column=c0, end_row=lay["band"], end_column=c1)
        for j, col in enumerate(cols):
            c = ws.cell(lay["head"], j + 1, label(col[1]))
            c.font, c.fill, c.alignment = BOLD, HEAD, WRAP
            if col[5]:
                c.comment = Comment(col[5], "model")
            ws.column_dimensions[get_column_letter(j + 1)].width = 12 if j > 2 else 7
        ws.row_dimensions[lay["head"]].height = 58
        for i in range(n_years):
            row = first + i
            for j, col in enumerate(cols):
                key, _, _, fn, fmt = col[:5]
                c = ws.cell(row, j + 1, fn(row, i))
                if fmt:
                    c.number_format = fmt
                if key in ("ret_override", "spend_mult"):
                    c.fill = INPUT
                if key in ("net_worth_today", "investments_today"):
                    c.fill = KEY
        ws.freeze_panes = "B1"
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 13
        ws.column_dimensions["C"].width = 13
        self.results_block()
        return ws

    def label(self, text):
        """A label that may mention the two people: becomes a formula so renaming them on household updates it."""
        if "{P1}" not in text and "{P2}" not in text:
            return text
        parts = text.replace("{P1}", "\x00P1\x00").replace("{P2}", "\x00P2\x00").split("\x00")
        bits = []
        for p in parts:
            if p == "P1":
                bits.append(self.P1)
            elif p == "P2":
                bits.append(self.P2)
            elif p:
                bits.append('"' + p.replace('"', '""') + '"')
        return "=" + "&".join(bits)

    def n_years(self):
        """Rows in the year table: until the younger person is 100 (at most 60)."""
        start = int(self.hh["start_year"])
        young = max(int(self.hh["p1_born"]), int(self.hh["p2_born"]))
        return max(10, min(60, young + 100 - start + 1))

    def property_formulas(self):
        ws, pc, TT, lay = self.ws, self.pc, self.TT, self.lay
        L, first, last = self.letters, self.first, self.last
        yr_rng = f"${L['year']}${first}:${L['year']}${last}"
        mb_rng = f"${L['mort_bal_start']}${first}:${L['mort_bal_start']}${last}"
        start = self.H["start_year"]
        for i in range(N_PROPS):
            pr, rr = 2 + i, lay["prop0"] + i
            B = f"$B{rr}"
            P = lambda k: f"properties!${pc[k]}${pr}"  # noqa: E731
            ws.cell(rr, 6, f'=IF({B}="",0,{P("value_now")}*(1+{P("appreciation")})^({B}-{start}+1))').number_format = MONEY
            net = f"F{rr}*(1-{P('selling_cost_pct')})"
            if i == 0:
                ws.cell(rr, 7, f'=IF({B}="",0,{net}-IFERROR(INDEX({mb_rng},MATCH({B},{yr_rng},0)),0))').number_format = MONEY
                ws.cell(rr, 8, f'=IF({B}="",0,MAX(0,{net}-{P("cost_basis")}-{TT["HOME_SALE_EXCLUSION"]}))' if self.us else 0).number_format = MONEY
                ws.cell(rr, 9, 0).number_format = MONEY
            else:
                taken = f"({P('accum_depreciation')}+{P('depreciation')}*MAX(0,MIN({B},{P('depreciation_end_year')})-{start}+1))"
                ws.cell(rr, 7, f'=IF({B}="",0,{net})').number_format = MONEY
                if self.us:
                    ws.cell(rr, 8, f'=IF({B}="",0,MAX(0,G{rr}-({P("cost_basis")}-{taken})))').number_format = MONEY
                    ws.cell(rr, 9, f'=IF({B}="",0,MIN(H{rr},{taken}))').number_format = MONEY
                else:
                    ws.cell(rr, 8, f'=IF({B}="",0,MAX(0,G{rr}-{P("cost_basis")}))').number_format = MONEY
                    ws.cell(rr, 9, f'=IF({B}="",0,MIN({taken},MAX(0,G{rr}-({P("cost_basis")}-{taken}))))').number_format = MONEY

    def model_columns(self):
        """(key, header, group, formula(row, i), number_format, comment)"""
        us, I, H, HC, TT, BR = self.us, self.I, self.H, self.HC, self.TT, self.BR
        c = lambda k, r: f"{self.letters[k]}{r}"          # noqa: E731
        p = lambda k, r: f"{self.letters[k]}{r - 1}"      # noqa: E731
        start = H["start_year"]
        p0, p1 = self.prop_rows
        e0, e1 = self.ev_rows
        blk = lambda col: f"${col}${p0}:${col}${p1}"                 # noqa: E731
        rblk = lambda col: f"${col}${p0 + 1}:${col}${p1}"            # noqa: E731
        prop = lambda key: f"properties!${self.pc[key]}$3:${self.pc[key]}${1 + N_PROPS}"   # noqa: E731
        home = lambda key: f"properties!${self.pc[key]}$2"         # noqa: E731
        active = lambda r: f"(1-({rblk('B')}>0)*({c('year', r)}>={rblk('B')}))"          # noqa: E731
        newuse = lambda r: f"(({rblk('C')}>0)*({c('year', r)}>={rblk('C')}))"             # noqa: E731
        ev = lambda t, r: (f"SUMPRODUCT(($B${e0}:$B${e1}={c('year', r)})*($D${e0}:$D${e1}=\"{t}\")*$C${e0}:$C${e1})"
                           f"*{c('infl_f', r)}")                                                     # noqa: E731

        def prog(x, key, idx=None):
            thr, _, step = BR[key]
            t = f"{thr}*{idx}" if idx else thr
            return f"SUMPRODUCT(({x}>{t})*({x}-{t})*{step})"

        def start_col(key, init):
            return lambda r, i: f"={I[init]}" if i == 0 else f"={p(key.replace('_start', '_end'), r)}"

        TAXF = lambda r: c("tax_f", r)       # noqa: E731
        F = lambda r: c("infl_f", r)         # noqa: E731
        cols = [
            ("year", "Year", "SUMMARY", lambda r, i: f"={start}+{i}" if i == 0 else f"={p('year', r)}+1", YEAR),
            ("age1", "{P1}'s age", "SUMMARY", lambda r, i: f"={c('year', r)}-{H['p1_born']}", "0"),
            ("age2", "{P2}'s age", "SUMMARY", lambda r, i: f"={c('year', r)}-{H['p2_born']}", "0"),
            ("living_costs", "Living costs (all spending except income tax)", "SUMMARY",
             lambda r, i: f"={c('housing_total', r)}+{c('household', r)}+{c('other_budget', r)}+{c('kids_parents', r)}+{c('college_shortfall', r)}+{c('health_total', r)}+{c('oneoff_expense', r)}", MONEY),
            ("taxes_paid", "Income + payroll tax paid this year", "SUMMARY",
             lambda r, i: f"={c('tax_true_up', r)}+{c('withholding', r)}+{c('payroll_tax', r)}", MONEY),
            ("gov_total", "Social Security" if us else "CPP + OAS", "SUMMARY",
             (lambda r, i: f"={c('ss1', r)}+{c('ss2', r)}") if us else
             (lambda r, i: f"={c('cpp1', r)}+{c('cpp2', r)}+{c('oas1', r)}+{c('oas2', r)}"), MONEY),
            ("draw", "Taken out of investments (net)", "SUMMARY",
             lambda r, i: (f"={c('sell_taxable', r)}+{c('from_pre', r)}+{c('from_free', r)}+{c('req1', r)}+{c('req2', r)}+{c('dividends', r)}-{c('surplus', r)}"
                           + (f"+{c('from_hsa', r)}+{c('hsa_used', r)}" if us else f"+{c('extra_pre', r)}+{c('to_free', r)}")), MONEY,
             "Everything withdrawn from your accounts this year (incl. required withdrawals and dividends), less anything reinvested"),
            ("draw_rate", "Withdrawal rate", "SUMMARY",
             lambda r, i: f"=IF({self._inv_start(c, r)}<=0,\"\",{c('draw', r)}/({self._inv_start(c, r)}))", PCT),
            ("investments_today", "Investments at year end (today's $)", "SUMMARY", lambda r, i: f"={c('investments_end', r)}/{F(r)}", MONEY),
            ("net_worth_today", "Net worth at year end (today's $)", "SUMMARY", lambda r, i: f"={c('net_worth_end', r)}/{F(r)}", MONEY,
             "Investments + home and property equity, in today's dollars. Before any tax due on death"),
            ("shortfall", "Shortfall (costs your accounts couldn't cover)", "SUMMARY",
             lambda r, i: (f"=MAX(0,{c('need', r)}-{c('sell_taxable', r)}-{c('from_pre', r)}-{c('from_free', r)}"
                           + (f"-{c('from_hsa', r)}" if us else "") + ")"), MONEY,
             "Non-zero means every account is empty - you'd have to sell property or cut spending"),
            # ------------------------------------------------ assumptions
            ("ret_override", "Return override (blank = default)", "ASSUMPTIONS BY YEAR", lambda r, i: None, PCT,
             "Type a return for any year to test a bad sequence, e.g. -20% then 0% for ten years"),
            ("ret", "Return used", "ASSUMPTIONS BY YEAR",
             lambda r, i: f"=IF({c('ret_override', r)}=\"\",{I['return_default']},{c('ret_override', r)})", PCT),
            ("spend_mult", "Household spending multiplier (blank = 100%)", "ASSUMPTIONS BY YEAR", lambda r, i: None, "0%",
             "Scale the household budget in any year, e.g. 90% from the year you downsize, 70% in a belt-tightening scenario"),
            ("infl_f", "Price level vs today", "ASSUMPTIONS BY YEAR", lambda r, i: f"=(1+{I['inflation']})^({c('year', r)}-{start}+1)", NUM3),
            ("tax_f", "Tax-law index", "ASSUMPTIONS BY YEAR", lambda r, i: f"=(1+{I['inflation']})^({c('year', r)}-{TT['TAX_YEAR']})", NUM3,
             "Price level vs TAX_YEAR: indexed tax amounts grow with it"),
        ]
        if not us:
            cols.append(("prov_f", "Provincial index", "ASSUMPTIONS BY YEAR",
                         lambda r, i: (f"=(1+{I['inflation']})^({c('year', r)}-{TT['TAX_YEAR']}-MAX(0,MIN({c('year', r)},"
                                       f"{TT['PROV_INDEX_PAUSE_TO']})-{TT['PROV_INDEX_PAUSE_FROM']}+1))"), NUM3,
                         "Tax-law index without the years the province isn't indexing"))
        cols += [
            ("med_f", "Healthcare price level", "ASSUMPTIONS BY YEAR", lambda r, i: f"=(1+{I['med_infl']})^({c('year', r)}-{start}+1)", NUM3),
            ("col_f", "College price level", "ASSUMPTIONS BY YEAR", lambda r, i: f"=(1+{I['col_infl']})^({c('year', r)}-{start}+1)", NUM3),
            ("ltc_f", "Long-term care price level", "ASSUMPTIONS BY YEAR", lambda r, i: f"=(1+{I['ltc_infl']})^({c('year', r)}-{start}+1)", NUM3),
            # ------------------------------------------------ work
            ("working1", "{P1} working?", "WORK", lambda r, i: f"=IF({c('year', r)}<{I['p1_retire']},1,0)", "0"),
            ("working2", "{P2} working?", "WORK", lambda r, i: f"=IF({c('year', r)}<{I['p2_retire']},1,0)", "0"),
            ("any_work", "Either working?", "WORK", lambda r, i: f"=IF({c('working1', r)}+{c('working2', r)}>0,1,0)", "0"),
            ("pay1", "{P1}'s pay", "WORK", lambda r, i: f"={c('working1', r)}*{I['p1_pay']}*{F(r)}+{ev(ONEOFF_TYPES[0], r)}", MONEY),
            ("pay2", "{P2}'s pay", "WORK", lambda r, i: f"={c('working2', r)}*{I['p2_pay']}*{F(r)}+{ev(ONEOFF_TYPES[1], r)}", MONEY),
            ("contrib1", "{P1}'s retirement contributions", "WORK", lambda r, i: f"={c('working1', r)}*{I['p1_contrib']}*{F(r)}", MONEY),
            ("contrib2", "{P2}'s retirement contributions", "WORK", lambda r, i: f"={c('working2', r)}*{I['p2_contrib']}*{F(r)}", MONEY),
            ("match1", "Employer contributions for {P1}", "WORK", lambda r, i: f"={c('working1', r)}*{I['p1_match']}*{F(r)}", MONEY),
            ("match2", "Employer contributions for {P2}", "WORK", lambda r, i: f"={c('working2', r)}*{I['p2_match']}*{F(r)}", MONEY),
        ]
        if us:
            cols.append(("hsa_contrib", "HSA contributions", "WORK", lambda r, i: f"={c('any_work', r)}*{I['hsa_contrib']}*{F(r)}", MONEY))
        cols += [
            ("gross_wages", "Total pay", "WORK", lambda r, i: f"={c('pay1', r)}+{c('pay2', r)}", MONEY),
            ("taxable_wages", "Taxable pay", "WORK",
             lambda r, i: f"={c('gross_wages', r)}-{c('contrib1', r)}-{c('contrib2', r)}" + (f"-{c('hsa_contrib', r)}" if us else ""), MONEY),
        ]
        if us:
            cols.append(("payroll_tax", "Social Security + Medicare payroll tax", "WORK",
                         lambda r, i: (f"=0.062*MIN({c('pay1', r)},{TT['SS_WAGE_BASE']}*{TAXF(r)})+0.062*MIN({c('pay2', r)},{TT['SS_WAGE_BASE']}*{TAXF(r)})"
                                       f"+0.0145*{c('gross_wages', r)}+0.009*MAX(0,{c('gross_wages', r)}-{TT['ADDITIONAL_MEDICARE_THRESHOLD']})"), MONEY))
        else:
            for k in (1, 2):
                cols.append((f"payroll{k}", f"{{P{k}}}'s CPP + EI contributions", "WORK",
                             (lambda k: lambda r, i: (f"={TT['CPP_RATE']}*MAX(0,MIN({c(f'pay{k}', r)},{TT['YMPE']}*{TAXF(r)})-{TT['CPP_EXEMPTION']})"
                                                      f"+{TT['CPP2_RATE']}*MAX(0,MIN({c(f'pay{k}', r)},{TT['YAMPE']}*{TAXF(r)})-{TT['YMPE']}*{TAXF(r)})"
                                                      f"+{TT['EI_RATE']}*MIN({c(f'pay{k}', r)},{TT['EI_MAX_INSURABLE']}*{TAXF(r)})"))(k), MONEY))
            cols.append(("payroll_tax", "CPP + EI contributions", "WORK", lambda r, i: f"={c('payroll1', r)}+{c('payroll2', r)}", MONEY))
        cols += [
            ("withholding", "Income tax withheld from pay", "WORK", lambda r, i: f"={I['withholding_rate']}*{c('taxable_wages', r)}", MONEY),
            ("net_pay", "Take-home pay", "WORK",
             lambda r, i: f"={c('gross_wages', r)}-{c('contrib1', r)}-{c('contrib2', r)}" + (f"-{c('hsa_contrib', r)}" if us else "")
             + f"-{c('payroll_tax', r)}-{c('withholding', r)}", MONEY),
            ("oneoff_taxable", "One-off: other taxable income", "WORK", lambda r, i: "=" + ev(ONEOFF_TYPES[2], r), MONEY),
            ("oneoff_taxfree", "One-off: tax-free money in", "WORK", lambda r, i: "=" + ev(ONEOFF_TYPES[3], r), MONEY),
            ("oneoff_expense", "One-off: expenses", "WORK", lambda r, i: "=" + ev(ONEOFF_TYPES[4], r), MONEY),
            # ------------------------------------------------ housing
            ("own_home", "Own your home?", "HOUSING", lambda r, i: f"=IF(OR($B${p0}=\"\",{c('year', r)}<$B${p0}),1,0)", "0"),
            ("mort_bal_start", "Mortgage balance, start", "HOUSING",
             lambda r, i: f"={home('mortgage_balance')}" if i == 0 else f"={p('mort_bal_end', r)}", MONEY),
            ("mort_interest", "Mortgage interest", "HOUSING", lambda r, i: f"={c('own_home', r)}*{c('mort_bal_start', r)}*{home('mortgage_rate')}", MONEY),
            ("mort_payment", "Mortgage payments", "HOUSING",
             lambda r, i: f"={c('own_home', r)}*MIN({home('mortgage_payment')},{c('mort_bal_start', r)}+{c('mort_interest', r)})", MONEY),
            ("mort_bal_end", "Mortgage balance, end", "HOUSING",
             lambda r, i: f"={c('own_home', r)}*({c('mort_bal_start', r)}+{c('mort_interest', r)}-{c('mort_payment', r)})", MONEY,
             "Paid off from the sale proceeds if you sell"),
            ("home_costs", "Home carrying costs", "HOUSING", lambda r, i: f"={c('own_home', r)}*{home('carrying_costs')}*{F(r)}", MONEY),
            ("new_housing", "Housing after selling your home", "HOUSING", lambda r, i: f"=(1-{c('own_home', r)})*$E${p0}*{F(r)}", MONEY),
            ("housing_total", "Housing total", "HOUSING",
             lambda r, i: f"={c('mort_payment', r)}+{c('home_costs', r)}+{c('new_housing', r)}", MONEY),
            # ------------------------------------------------ living
            ("household", "Household living (budget tab)", "LIVING",
             lambda r, i: (f"=SUMIF({self.budget_rng[0]},\"household\",{self.budget_rng[1]})*{F(r)}"
                           f"*IF({c('spend_mult', r)}=\"\",1,{c('spend_mult', r)})"), MONEY),
            ("other_budget", "Renovations, cars, other (other_budget tab)", "LIVING",
             lambda r, i: f"=SUMPRODUCT(({self.ob['start']}<={c('year', r)})*({self.ob['end']}>={c('year', r)})*{self.ob['amt']})*{F(r)}", MONEY),
            ("kids_parents", "Kids: costs you pay", "LIVING",
             lambda r, i: (f"=SUMPRODUCT(({self.kb_start}<={c('year', r)})*({self.kb_end}>={c('year', r)})*({self.kb_paid}<>\"college_funds\")*{self.kb_amt}"
                           f"*(({self.kb_infl}=\"college\")*{c('col_f', r)}+({self.kb_infl}<>\"college\")*{F(r)}))"), MONEY),
            ("college_cost", "Kids: costs paid from college funds", "LIVING",
             lambda r, i: (f"=SUMPRODUCT(({self.kb_start}<={c('year', r)})*({self.kb_end}>={c('year', r)})*({self.kb_paid}=\"college_funds\")*{self.kb_amt}"
                           f"*(({self.kb_infl}=\"college\")*{c('col_f', r)}+({self.kb_infl}<>\"college\")*{F(r)}))"), MONEY),
            ("college_start", "College funds, start", "LIVING", start_col("college_start", "college_0"), MONEY),
            ("college_from_funds", "Paid by college funds", "LIVING",
             lambda r, i: f"=MIN({c('college_cost', r)},{c('college_start', r)}*(1+{c('ret', r)}))", MONEY),
            ("college_shortfall", "College costs paid by you", "LIVING", lambda r, i: f"={c('college_cost', r)}-{c('college_from_funds', r)}", MONEY),
            ("college_end", "College funds, end", "LIVING",
             lambda r, i: f"={c('college_start', r)}*(1+{c('ret', r)})-{c('college_from_funds', r)}", MONEY, "What's left stays the kids' money"),
            # ------------------------------------------------ healthcare
            ("employer_health", "Employer health plan (while working)", "HEALTHCARE",
             lambda r, i: f"={c('any_work', r)}*{HC['employer_plan_cost']}*{c('med_f', r)}", MONEY),
        ]
        if us:
            cols += [
                ("adults_marketplace", "Adults buying their own insurance", "HEALTHCARE",
                 lambda r, i: f"=(1-{c('any_work', r)})*(({c('age1', r)}<65)+({c('age2', r)}<65))", "0"),
                ("kids_covered", "Kids on your plan", "HEALTHCARE", lambda r, i: f"=COUNTIF({H['kids_through']},\">=\"&{c('year', r)})", "0"),
                ("marketplace_premium", "Marketplace premiums", "HEALTHCARE",
                 lambda r, i: (f"=({c('adults_marketplace', r)}*{HC['marketplace_adult_monthly']}+IF({c('adults_marketplace', r)}>0,{c('kids_covered', r)},0)"
                               f"*{HC['marketplace_child_monthly']})*12*{c('med_f', r)}"), MONEY, "Unsubsidized"),
                ("marketplace_oop", "Marketplace out-of-pocket", "HEALTHCARE",
                 lambda r, i: f"=IF({c('adults_marketplace', r)}>0,({c('adults_marketplace', r)}+{c('kids_covered', r)})*{HC['marketplace_oop_per_person']}*{c('med_f', r)},0)", MONEY),
                ("medicare_n", "People on Medicare", "HEALTHCARE", lambda r, i: f"=({c('age1', r)}>=65)+({c('age2', r)}>=65)", "0"),
                ("medicare_premiums", "Medicare B + D + Medigap premiums", "HEALTHCARE",
                 lambda r, i: f"={c('medicare_n', r)}*({HC['medicare_part_b_monthly']}+{HC['medicare_part_d_monthly']}+{HC['medigap_monthly']})*12*{c('med_f', r)}", MONEY),
                ("medicare_oop", "Medicare out-of-pocket", "HEALTHCARE",
                 lambda r, i: f"={c('medicare_n', r)}*{HC['medicare_oop_per_person']}*{c('med_f', r)}", MONEY),
                ("irmaa", "Medicare income surcharge (IRMAA)", "HEALTHCARE",
                 lambda r, i: "=0" if i < 2 else (f"={c('medicare_n', r)}*12*{F(r)}*SUMPRODUCT(({self.letters['agi']}{r - 2}>{BR['irmaa'][0]}*{TAXF(r)})*{BR['irmaa'][1]})"),
                 MONEY, "Based on income two years earlier"),
            ]
        else:
            cols += [
                ("retired_n", "Adults on a retiree health plan", "HEALTHCARE", lambda r, i: f"=(1-{c('any_work', r)})*2", "0"),
                ("private_premium", "Extended health + dental premiums", "HEALTHCARE",
                 lambda r, i: f"={c('retired_n', r)}*{HC['private_plan_monthly']}*12*{c('med_f', r)}", MONEY, "Provincial health insurance covers the rest"),
                ("health_oop", "Out-of-pocket (drugs, dental, vision)", "HEALTHCARE",
                 lambda r, i: f"={c('retired_n', r)}*{HC['oop_per_person']}*{c('med_f', r)}", MONEY),
            ]
        cols += [
            ("ltc_people", "People in long-term care", "HEALTHCARE",
             lambda r, i: (f"=AND({c('age1', r)}>={I['p1_ltc_age']},{c('age1', r)}<{I['p1_ltc_age']}+{I['p1_ltc_years']})*1"
                           f"+AND({c('age2', r)}>={I['p2_ltc_age']},{c('age2', r)}<{I['p2_ltc_age']}+{I['p2_ltc_years']})*1"), "0"),
            ("ltc_cost", "Long-term care", "HEALTHCARE", lambda r, i: f"={c('ltc_people', r)}*{I['ltc_cost']}*{c('ltc_f', r)}", MONEY),
            ("health_total", "Healthcare total", "HEALTHCARE",
             (lambda r, i: f"={c('employer_health', r)}+{c('marketplace_premium', r)}+{c('marketplace_oop', r)}+{c('medicare_premiums', r)}+{c('medicare_oop', r)}+{c('irmaa', r)}+{c('ltc_cost', r)}")
             if us else (lambda r, i: f"={c('employer_health', r)}+{c('private_premium', r)}+{c('health_oop', r)}+{c('ltc_cost', r)}"), MONEY),
        ]
        if us:
            cols += [
                ("hsa_start", "HSA, start", "HEALTHCARE", start_col("hsa_start", "hsa_0"), MONEY),
                ("hsa_used", "Paid from HSA", "HEALTHCARE",
                 lambda r, i: (f"=IF({c('any_work', r)}=1,0,MIN({c('hsa_start', r)}*(1+{c('ret', r)}),{c('marketplace_oop', r)}"
                               f"+{c('medicare_n', r)}*({HC['medicare_part_b_monthly']}+{HC['medicare_part_d_monthly']})*12*{c('med_f', r)}"
                               f"+{c('medicare_oop', r)}+{c('irmaa', r)}+{c('ltc_cost', r)}))"), MONEY,
                 "Out-of-pocket, Medicare B/D premiums, IRMAA and long-term care are HSA-eligible"),
                ("hsa_end", "HSA, end", "HEALTHCARE",
                 lambda r, i: f"={c('hsa_start', r)}*(1+{c('ret', r)})-{c('hsa_used', r)}+{c('hsa_contrib', r)}-{c('from_hsa', r)}", MONEY),
            ]
        # ------------------------------------------------ properties
        cols += [
            ("rental_cash", "Property income after costs", "PROPERTIES",
             lambda r, i: (f"=SUMPRODUCT({active(r)}*(1-{newuse(r)})*({prop('rent')}-{prop('cash_costs')}))*{F(r)}"
                           f"+SUMPRODUCT({active(r)}*{newuse(r)}*({rblk('D')}-{rblk('E')}))*{F(r)}"), MONEY,
             "Other properties' rent less all their cash costs, unless sold or changed on this tab"),
            ("rental_dep", "Depreciation" if us else "Capital cost allowance", "PROPERTIES",
             lambda r, i: f"=SUMPRODUCT({active(r)}*{prop('depreciation')}*({c('year', r)}<={prop('depreciation_end_year')}))", MONEY),
            ("rental_taxable", "Taxable property income", "PROPERTIES", lambda r, i: f"=MAX(0,{c('rental_cash', r)}-{c('rental_dep', r)})", MONEY,
             "Losses aren't deducted (conservative)"),
            ("sale_proceeds", "Property sale proceeds", "PROPERTIES", lambda r, i: f"=SUMPRODUCT(({blk('B')}={c('year', r)})*{blk('G')})", MONEY),
            ("sale_gain", "Taxable gain on sales" if us else "Capital gains on sales", "PROPERTIES",
             lambda r, i: f"=SUMPRODUCT(({blk('B')}={c('year', r)})*{blk('H')})", MONEY),
            ("sale_recapture", "Depreciation recapture" if us else "CCA recapture", "PROPERTIES",
             lambda r, i: f"=SUMPRODUCT(({blk('B')}={c('year', r)})*{blk('I')})", MONEY),
            ("home_value", "Home value, end", "PROPERTIES",
             lambda r, i: f"={c('own_home', r)}*{home('value_now')}*(1+{home('appreciation')})^({c('year', r)}-{start}+1)", MONEY),
            ("rental_value", "Other property values, end", "PROPERTIES",
             lambda r, i: f"=SUMPRODUCT({active(r)}*{prop('value_now')}*(1+{prop('appreciation')})^({c('year', r)}-{start}+1))", MONEY),
        ]
        # ------------------------------------------------ government pensions
        if us:
            cut = lambda r: f"IF({c('year', r)}>={I['ss_cut_year']},1-{I['ss_cut_pct']},1)"  # noqa: E731
            own = lambda k, r: f"{H[f'p{k}_ss_benefit']}*12*(1-{I['ss_early_cut']})*{F(r)}*{cut(r)}"  # noqa: E731
            for k in (1, 2):
                j = 3 - k
                cols.append((f"ss{k}", f"{{P{k}}}'s Social Security", "SOCIAL SECURITY",
                             (lambda k, j: lambda r, i: (f"=IF({c('year', r)}>={I[f'p{k}_claim_year']},IF({c('year', r)}>={I[f'p{j}_claim_year']},"
                                                         f"MAX({own(k, r)}*{I[f'p{k}_ss_factor']},0.5*{own(j, r)}*{I[f'p{k}_spousal_factor']}),"
                                                         f"{own(k, r)}*{I[f'p{k}_ss_factor']}),0)"))(k, j), MONEY))
        else:
            for k in (1, 2):
                cols.append((f"cpp{k}", f"{{P{k}}}'s CPP", "CPP AND OAS",
                             (lambda k: lambda r, i: (f"=IF({c('year', r)}>={H[f'p{k}_born']}+{I[f'p{k}_cpp_age']},"
                                                      f"{H[f'p{k}_cpp_at_65']}*12*{I[f'p{k}_cpp_factor']}*{F(r)},0)"))(k), MONEY))
            for k in (1, 2):
                cols.append((f"oas{k}", f"{{P{k}}}'s OAS", "CPP AND OAS",
                             (lambda k: lambda r, i: (f"=IF({c('year', r)}>={H[f'p{k}_born']}+{I[f'p{k}_oas_age']},{TT['OAS_MONTHLY']}*12*{TAXF(r)}"
                                                      f"*MIN(1,{H[f'p{k}_oas_years']}/40)*{I[f'p{k}_oas_factor']}*(1+{TT['OAS_75_BOOST']}*({c(f'age{k}', r)}>=75)),0)"))(k),
                             MONEY, "Before the clawback (that's in the tax columns)"))
        for k in (1, 2):
            cols.append((f"pension{k}", f"{{P{k}}}'s workplace pension", "PENSIONS",
                         (lambda k: lambda r, i: (
                             f"=IF(AND({I[f'p{k}_pension']}<>0,{c('year', r)}>={H[f'p{k}_born']}+{I[f'p{k}_pension_age']}),{I[f'p{k}_pension']}"
                             f"*(1+{I['inflation']})^(MAX({H[f'p{k}_born']}+{I[f'p{k}_pension_age']},{start}-1)-{start}+1)"
                             f"*(1+{I['pension_indexing']}*{I['inflation']})^({c('year', r)}-MAX({H[f'p{k}_born']}+{I[f'p{k}_pension_age']},{start}-1)),0)"))(k),
                         MONEY))
        # ------------------------------------------------ investments
        pre_name = "pre-tax" if us else "RRSP/RRIF"
        free_name = "Roth" if us else "TFSA"
        tax_name = "Taxable" if us else "Non-registered"
        cols += [
            ("taxable_start", f"{tax_name}, start", "INVESTMENTS", start_col("taxable_start", "taxable_0"), MONEY),
            ("basis_start", "Cost basis, start" if us else "Adjusted cost base, start", "INVESTMENTS", start_col("basis_start", "basis_0"), MONEY),
            ("dividends", "Dividends", "INVESTMENTS", lambda r, i: f"={I['dividend_yield']}*{c('taxable_start', r)}", MONEY),
            ("pre1_start", f"{{P1}} {pre_name}, start", "INVESTMENTS", start_col("pre1_start", "pre1_0"), MONEY),
            ("pre2_start", f"{{P2}} {pre_name}, start", "INVESTMENTS", start_col("pre2_start", "pre2_0"), MONEY),
        ]
        for k in (1, 2):
            if us:
                f_ = (lambda k: lambda r, i: (f"=IF({c(f'age{k}', r)}>={I[f'p{k}_rmd_age']},MAX(0,{c(f'pre{k}_start', r)})"
                                              f"/INDEX({BR['factors'][1]},MATCH({c(f'age{k}', r)},{BR['factors'][0]},0)),0)"))(k)
            else:
                f_ = (lambda k: lambda r, i: (f"=IF({c(f'age{k}', r)}>={TT['RRIF_START_AGE']},MAX(0,{c(f'pre{k}_start', r)})"
                                              f"*INDEX({BR['factors'][1]},MATCH(MIN({c(f'age{k}', r)}-1,120),{BR['factors'][0]},0)),0)"))(k)
            cols.append((f"req{k}", f"{{P{k}}}'s required withdrawal" if us else f"{{P{k}}}'s RRIF minimum", "INVESTMENTS", f_, MONEY))
        cols += [
            ("free_start", f"{free_name}, start", "INVESTMENTS", start_col("free_start", "free_0"), MONEY),
            ("conv_want", "Roth conversion planned" if us else "Extra RRSP/RRIF withdrawal planned", "INVESTMENTS",
             lambda r, i: (f"=IF(OR({I['conv_from']}=\"\",{I['conv_to']}=\"\"),0,IF(AND({c('year', r)}>={I['conv_from']},{c('year', r)}<={I['conv_to']}),"
                           f"{I['conv_amount']}*{F(r)},0))"), MONEY),
            ("avail1", f"{{P1}} {pre_name} available", "INVESTMENTS", lambda r, i: f"=MAX(0,{c('pre1_start', r)}*(1+{c('ret', r)})-{c('req1', r)})", MONEY),
            ("avail2", f"{{P2}} {pre_name} available", "INVESTMENTS", lambda r, i: f"=MAX(0,{c('pre2_start', r)}*(1+{c('ret', r)})-{c('req2', r)})", MONEY),
        ]
        if us:
            left = lambda k, r: c(f"avail{k}", r)  # noqa: E731
        else:
            cols += [
                ("extra_pre", "Extra RRSP/RRIF withdrawal", "INVESTMENTS",
                 lambda r, i: f"=MIN({c('conv_want', r)},{c('avail1', r)}+{c('avail2', r)})", MONEY, "Taken as cash; taxed as income"),
                ("melt1", "... from {P1}'s", "INVESTMENTS",
                 lambda r, i: f"=IF({c('avail1', r)}+{c('avail2', r)}>1,{c('extra_pre', r)}*{c('avail1', r)}/({c('avail1', r)}+{c('avail2', r)}),0)", MONEY),
                ("melt2", "... from {P2}'s", "INVESTMENTS", lambda r, i: f"={c('extra_pre', r)}-{c('melt1', r)}", MONEY),
            ]
            left = lambda k, r: f"({c(f'avail{k}', r)}-{c(f'melt{k}', r)})"  # noqa: E731
        if not us:
            cols.append(("tfsa_room", "TFSA room, start", "INVESTMENTS",
                         lambda r, i: f"={H['tfsa_room']}" if i == 0 else
                         f"={p('tfsa_room', r)}-{p('to_free', r)}+{p('from_free', r)}+{TT['TFSA_LIMIT']}*{TAXF(r)}*2", MONEY))
        cols += [
            ("tax_true_up", "Last year's income tax still owed", "INVESTMENTS",
             lambda r, i: f"={I['tax_due_first']}" if i == 0 else f"={p('income_tax', r)}-{p('withholding', r)}", MONEY,
             "Taxes are settled the following April; negative = refund"),
            ("cash_in", "Cash coming in", "INVESTMENTS",
             lambda r, i: (f"={c('net_pay', r)}+{c('gov_total', r)}+{c('pension1', r)}+{c('pension2', r)}+{c('rental_cash', r)}+{c('dividends', r)}+{c('req1', r)}+{c('req2', r)}"
                           f"+{c('sale_proceeds', r)}+{c('oneoff_taxable', r)}+{c('oneoff_taxfree', r)}"
                           + (f"+{c('hsa_used', r)}" if us else f"+{c('extra_pre', r)}")), MONEY),
            ("need", "Still needed (+) or surplus (-)", "INVESTMENTS", lambda r, i: f"={c('living_costs', r)}+{c('tax_true_up', r)}-{c('cash_in', r)}", MONEY),
            ("sell_taxable", f"Sold from {tax_name.lower()}", "INVESTMENTS",
             lambda r, i: f"=MAX(0,MIN({c('need', r)},{c('taxable_start', r)}*(1+{c('ret', r)}-{I['dividend_yield']})))", MONEY,
             f"Order: {tax_name.lower()} first, then {pre_name} (split by what each of you has), then {free_name}" + (", then HSA" if us else "")),
            ("from_pre", f"Withdrawn from {pre_name}", "INVESTMENTS",
             lambda r, i: f"=MAX(0,MIN({c('need', r)}-{c('sell_taxable', r)},{left(1, r)}+{left(2, r)}))", MONEY),
            ("w1", "... from {P1}'s", "INVESTMENTS",
             lambda r, i: f"=IF({left(1, r)}+{left(2, r)}>1,{c('from_pre', r)}*{left(1, r)}/({left(1, r)}+{left(2, r)}),0)", MONEY),
            ("w2", "... from {P2}'s", "INVESTMENTS", lambda r, i: f"={c('from_pre', r)}-{c('w1', r)}", MONEY),
            ("from_free", f"Withdrawn from {free_name}", "INVESTMENTS",
             lambda r, i: f"=MAX(0,MIN({c('need', r)}-{c('sell_taxable', r)}-{c('from_pre', r)},{c('free_start', r)}*(1+{c('ret', r)})))", MONEY),
        ]
        if us:
            cols += [
                ("from_hsa", "Extra from HSA (last resort)", "INVESTMENTS",
                 lambda r, i: (f"=MAX(0,MIN({c('need', r)}-{c('sell_taxable', r)}-{c('from_pre', r)}-{c('from_free', r)},"
                               f"{c('hsa_start', r)}*(1+{c('ret', r)})-{c('hsa_used', r)}+{c('hsa_contrib', r)}))"), MONEY),
            ]
        cols += [("surplus", "Surplus reinvested", "INVESTMENTS", lambda r, i: f"=MAX(0,-{c('need', r)})", MONEY)]
        if us:
            cols += [
                ("conversion", "Roth conversion", "INVESTMENTS",
                 lambda r, i: f"=MIN({c('conv_want', r)},{c('avail1', r)}-{c('w1', r)}+{c('avail2', r)}-{c('w2', r)})", MONEY,
                 "Pre-tax money left after withdrawals, moved to the Roth"),
                ("conv1", "... from {P1}'s", "INVESTMENTS",
                 lambda r, i: (f"=IF({c('avail1', r)}-{c('w1', r)}+{c('avail2', r)}-{c('w2', r)}>1,{c('conversion', r)}*({c('avail1', r)}-{c('w1', r)})"
                               f"/({c('avail1', r)}-{c('w1', r)}+{c('avail2', r)}-{c('w2', r)}),0)"), MONEY),
                ("conv2", "... from {P2}'s", "INVESTMENTS", lambda r, i: f"={c('conversion', r)}-{c('conv1', r)}", MONEY),
            ]
            to_tax = lambda r: c("surplus", r)  # noqa: E731
        else:
            cols += [("to_free", "Surplus into TFSAs", "INVESTMENTS", lambda r, i: f"=MIN({c('surplus', r)},MAX(0,{c('tfsa_room', r)}))", MONEY,
                      "Up to the TFSA room; the rest goes to the non-registered account")]
            to_tax = lambda r: f"({c('surplus', r)}-{c('to_free', r)})"  # noqa: E731
        cols += [
            ("gain_frac", "Gain share of sales", "INVESTMENTS",
             lambda r, i: f"=IF({c('taxable_start', r)}>0,MAX(0,1-{c('basis_start', r)}/{c('taxable_start', r)}),0)", PCT),
            ("realized_gain", "Capital gains realized", "INVESTMENTS", lambda r, i: f"={c('sell_taxable', r)}*{c('gain_frac', r)}", MONEY),
            ("taxable_end", f"{tax_name}, end", "INVESTMENTS",
             lambda r, i: f"={c('taxable_start', r)}*(1+{c('ret', r)}-{I['dividend_yield']})-{c('sell_taxable', r)}+{to_tax(r)}", MONEY),
            ("basis_end", "Cost basis, end" if us else "Adjusted cost base, end", "INVESTMENTS",
             lambda r, i: f"={c('basis_start', r)}-({c('sell_taxable', r)}-{c('realized_gain', r)})+{to_tax(r)}", MONEY),
        ]
        for k in (1, 2):
            cols.append((f"pre{k}_end", f"{{P{k}}} {pre_name}, end", "INVESTMENTS",
                         (lambda k: lambda r, i: (f"={c(f'pre{k}_start', r)}*(1+{c('ret', r)})-{c(f'req{k}', r)}-{c(f'w{k}', r)}"
                                                  + (f"-{c(f'conv{k}', r)}" if us else f"-{c(f'melt{k}', r)}")
                                                  + f"+{c(f'contrib{k}', r)}+{c(f'match{k}', r)}"))(k), MONEY))
        cols.append(("free_end", f"{free_name}, end", "INVESTMENTS",
                     lambda r, i: f"={c('free_start', r)}*(1+{c('ret', r)})-{c('from_free', r)}+" + (c("conversion", r) if us else c("to_free", r)), MONEY))
        # ------------------------------------------------ taxes
        if us:
            cols += self._us_tax_columns(c, p, TAXF)
        else:
            cols += self._ca_tax_columns(c, p, TAXF)
        cols += [
            ("investments_end", "Investments, end", "TOTALS",
             lambda r, i: f"={c('taxable_end', r)}+{c('pre1_end', r)}+{c('pre2_end', r)}+{c('free_end', r)}" + (f"+{c('hsa_end', r)}" if us else ""), MONEY,
             "Your accounts; excludes the kids' college funds"),
            ("net_worth_end", "Net worth, end", "TOTALS",
             lambda r, i: f"={c('investments_end', r)}+{c('home_value', r)}-{c('mort_bal_end', r)}+{c('rental_value', r)}", MONEY),
        ]
        return [col if len(col) == 6 else (*col, "") for col in cols]

    def _inv_start(self, c, r):
        s = f"{c('taxable_start', r)}+{c('pre1_start', r)}+{c('pre2_start', r)}+{c('free_start', r)}"
        return s + (f"+{c('hsa_start', r)}" if self.us else "")

    def _us_tax_columns(self, c, p, TAXF):
        I, TT, BR = self.I, self.TT, self.BR

        def prog(x, key, idx):
            thr, _, step = BR[key]
            return f"SUMPRODUCT(({x}>{thr}*{idx})*({x}-{thr}*{idx})*{step})"
        ss = lambda r: c("gov_total", r)  # noqa: E731
        prov = lambda r: f"({c('ordinary', r)}+{c('pref', r)}+0.5*{ss(r)})"  # noqa: E731
        sidx = lambda r: f"IF({TT['STATE_INDEXED']}=1,{TAXF(r)},1)"  # noqa: E731
        return [
            ("ordinary", "Ordinary income", "INCOME TAX",
             lambda r, i: (f"={c('taxable_wages', r)}+{c('from_pre', r)}+{c('req1', r)}+{c('req2', r)}+{c('conversion', r)}+{c('from_hsa', r)}"
                           f"+{c('rental_taxable', r)}+{c('oneoff_taxable', r)}+{c('pension1', r)}+{c('pension2', r)}"), MONEY),
            ("pref", "Dividends + capital gains", "INCOME TAX", lambda r, i: f"={c('dividends', r)}+{c('realized_gain', r)}+{c('sale_gain', r)}", MONEY),
            ("ss_taxable", "Taxable Social Security", "INCOME TAX",
             lambda r, i: (f"=IF({prov(r)}<={TT['SS_TAXABLE_BASE']},0,IF({prov(r)}<={TT['SS_TAXABLE_ADJUSTED_BASE']},"
                           f"MIN(0.5*{ss(r)},0.5*({prov(r)}-{TT['SS_TAXABLE_BASE']})),"
                           f"MIN(0.85*{ss(r)},0.85*({prov(r)}-{TT['SS_TAXABLE_ADJUSTED_BASE']})+MIN(0.5*{ss(r)},{TT['SS_TAXABLE_ADJUSTED_BASE']}-{TT['SS_TAXABLE_BASE']}))))"), MONEY),
            ("agi", "Adjusted gross income", "INCOME TAX", lambda r, i: f"={c('ordinary', r)}+{c('pref', r)}+{c('ss_taxable', r)}", MONEY),
            ("n65", "People 65+", "INCOME TAX", lambda r, i: f"=({c('age1', r)}>=65)+({c('age2', r)}>=65)", "0"),
            ("std_deduction", "Standard deduction", "INCOME TAX",
             lambda r, i: f"=({TT['FED_STD']}+{TT['FED_65_ADD']}*{c('n65', r)})*{TAXF(r)}", MONEY),
            ("senior_deduction", "Senior deduction", "INCOME TAX",
             lambda r, i: (f"=IF(AND({c('year', r)}>={TT['SENIOR_DED_FROM']},{c('year', r)}<={TT['SENIOR_DED_TO']}),{c('n65', r)}"
                           f"*MAX(0,{TT['SENIOR_DED']}-{TT['SENIOR_PHASEOUT_RATE']}*MAX(0,{c('agi', r)}-{TT['SENIOR_PHASEOUT_START']})),0)"), MONEY,
             "$6,000 per person 65+ in 2025-2028, on top of the standard or itemized deduction"),
            ("salt_cap", "SALT cap", "INCOME TAX",
             lambda r, i: (f"=IFERROR(MAX({TT['SALT_CAP_AFTER']},INDEX({BR['salt'][1]},MATCH({c('year', r)},{BR['salt'][0]},0))"
                           f"-{TT['SALT_PHASEDOWN_RATE']}*MAX(0,{c('agi', r)}-{TT['SALT_PHASEDOWN_START']}*1.01^({c('year', r)}-{TT['TAX_YEAR']}))),{TT['SALT_CAP_AFTER']})"), MONEY),
            ("medical_deduction", "Medical deduction", "INCOME TAX",
             lambda r, i: f"=MAX(0,{c('health_total', r)}-{c('employer_health', r)}-{TT['MEDICAL_FLOOR']}*{c('agi', r)})", MONEY,
             "Premiums, out-of-pocket and long-term care above 7.5% of AGI"),
            ("itemized", "Itemized deductions", "INCOME TAX",
             lambda r, i: (f"={c('mort_interest', r)}+MIN({c('salt_cap', r)},"
                           + ("0" if i == 0 else f"{p('state_tax', r)}+{p('local_tax', r)}+{p('other_state_tax', r)}")
                           + f"+{c('own_home', r)}*properties!${self.pc['property_tax']}$2*{c('infl_f', r)})+{c('medical_deduction', r)}+{I['charity']}*{c('infl_f', r)}"), MONEY,
             "Mortgage interest, state/local and property taxes (capped), medical, charity"),
            ("deduction", "Deduction used", "INCOME TAX", lambda r, i: f"=MAX({c('std_deduction', r)},{c('itemized', r)})", MONEY),
            ("ord_ti", "Taxable ordinary income", "INCOME TAX",
             lambda r, i: f"=MAX(0,{c('ordinary', r)}+{c('ss_taxable', r)}-{c('deduction', r)}-{c('senior_deduction', r)})", MONEY),
            ("pref_ti", "Taxable gains + dividends", "INCOME TAX",
             lambda r, i: f"=MAX(0,{c('pref', r)}-MAX(0,{c('deduction', r)}+{c('senior_deduction', r)}-{c('ordinary', r)}-{c('ss_taxable', r)}))", MONEY),
            ("fed_tax", "Federal income tax", "INCOME TAX",
             lambda r, i: ("=" + prog(c('ord_ti', r), "fed", TAXF(r))
                           + "+" + prog(f"({c('ord_ti', r)}+{c('pref_ti', r)})", "ltcg", TAXF(r))
                           + "-" + prog(c('ord_ti', r), "ltcg", TAXF(r))
                           + f"+{TT['UNRECAPTURED_1250_EXTRA']}*{c('sale_recapture', r)}"
                           + f"+{TT['NIIT_RATE']}*MIN({c('dividends', r)}+{c('realized_gain', r)}+{c('rental_taxable', r)}+{c('sale_gain', r)},MAX(0,{c('agi', r)}-{TT['NIIT_THRESHOLD']}))"), MONEY,
             "Brackets + gains stacked on top + 3.8% investment tax + recapture"),
            ("moved", "Moved to another state?", "INCOME TAX",
             lambda r, i: f"=IF({I['move_year']}=\"\",0,IF({c('year', r)}>={I['move_year']},1,0))", "0"),
            ("ret_exclusion", "State retirement-income exclusion", "INCOME TAX",
             lambda r, i: (f"=IF({c('age1', r)}+0.5>={TT['STATE_RET_EXCLUSION_AGE']},MIN({c('w1', r)}+{c('req1', r)}+{c('conv1', r)}+{c('pension1', r)},{TT['STATE_RET_EXCLUSION']}),0)"
                           f"+IF({c('age2', r)}+0.5>={TT['STATE_RET_EXCLUSION_AGE']},MIN({c('w2', r)}+{c('req2', r)}+{c('conv2', r)}+{c('pension2', r)},{TT['STATE_RET_EXCLUSION']}),0)"), MONEY),
            ("state_taxable", "State taxable income", "INCOME TAX",
             lambda r, i: (f"=MAX(0,{c('agi', r)}-{TT['STATE_SS_EXEMPT']}*{c('ss_taxable', r)}-{c('ret_exclusion', r)}-{TT['STATE_STD']}*{sidx(r)})"), MONEY),
            ("state_tax", "State income tax", "INCOME TAX", lambda r, i: f"=(1-{c('moved', r)})*" + prog(c('state_taxable', r), "state", sidx(r)), MONEY),
            ("local_tax", "Local income tax", "INCOME TAX", lambda r, i: f"=(1-{c('moved', r)})*{TT['LOCAL_RATE']}*{c('state_taxable', r)}", MONEY),
            ("other_state_tax", "New state's income tax", "INCOME TAX",
             lambda r, i: f"={c('moved', r)}*{I['other_state_rate']}*MAX(0,{c('agi', r)}-{c('ss_taxable', r)})", MONEY),
            ("income_tax", "Income tax for the year", "INCOME TAX",
             lambda r, i: f"={c('fed_tax', r)}+{c('state_tax', r)}+{c('local_tax', r)}+{c('other_state_tax', r)}", MONEY),
        ]

    def _ca_tax_columns(self, c, p, TAXF):
        I, H, TT, BR = self.I, self.H, self.TT, self.BR
        PF = lambda r: c("prov_f", r)  # noqa: E731

        def prog(x, key, idx):
            thr, _, step = BR[key]
            return f"SUMPRODUCT(({x}>{thr}*{idx})*({x}-{thr}*{idx})*{step})"
        share = lambda k: H["joint_share_p1"] if k == 1 else f"(1-{H['joint_share_p1']})"  # noqa: E731
        hi = lambda r: f"({c('pre_split1', r)}>={c('pre_split2', r)})"  # noqa: E731
        signed = lambda k, r: (f"IF({hi(r)},-{c('split', r)},{c('split', r)})" if k == 1 else
                               f"IF({hi(r)},{c('split', r)},-{c('split', r)})")  # noqa: E731
        cols = []
        for k in (1, 2):
            cols += [
                (f"rrif{k}", f"{{P{k}}}'s RRSP/RRIF income", "INCOME TAX",
                 (lambda k: lambda r, i: f"={c(f'req{k}', r)}+{c(f'melt{k}', r)}+{c(f'w{k}', r)}")(k), MONEY),
                (f"gross_div{k}", f"{{P{k}}}'s grossed-up eligible dividends", "INCOME TAX",
                 (lambda k: lambda r, i: f"={c('dividends', r)}*{share(k)}*{I['eligible_div_share']}*(1+{TT['ELIGIBLE_GROSSUP']})")(k), MONEY),
                (f"pre_split{k}", f"{{P{k}}}'s net income before pension splitting", "INCOME TAX",
                 (lambda k: lambda r, i: (f"={c(f'pay{k}', r)}+{c(f'cpp{k}', r)}+{c(f'oas{k}', r)}+{c(f'pension{k}', r)}+{c(f'rrif{k}', r)}+{c(f'gross_div{k}', r)}"
                                          f"+{c('dividends', r)}*{share(k)}*(1-{I['eligible_div_share']})"
                                          f"+({c('realized_gain', r)}+{c('sale_gain', r)})*{share(k)}*{TT['CG_INCLUSION']}"
                                          f"+({c('sale_recapture', r)}+{c('rental_taxable', r)}+{c('oneoff_taxable', r)})*{share(k)}-{c(f'contrib{k}', r)}"))(k), MONEY,
                 "Joint investment and property income is split by the household tab's share"),
            ]
        cols.append(("split", "Pension income moved to the lower-income spouse", "INCOME TAX",
                      lambda r, i: (f"=IF({hi(r)},MIN(0.5*({c('pension1', r)}+IF({c('age1', r)}>=65,{c('rrif1', r)},0)),MAX(0,({c('pre_split1', r)}-{c('pre_split2', r)})/2)),"
                                    f"MIN(0.5*({c('pension2', r)}+IF({c('age2', r)}>=65,{c('rrif2', r)},0)),MAX(0,({c('pre_split2', r)}-{c('pre_split1', r)})/2)))"), MONEY,
                      "Workplace pensions (any age) and RRIF income (65+) can be split up to 50%; the model evens out your incomes"))
        for k in (1, 2):
            cols += [
                (f"net_income{k}", f"{{P{k}}}'s net income", "INCOME TAX", (lambda k: lambda r, i: f"={c(f'pre_split{k}', r)}+{signed(k, r)}")(k), MONEY),
                (f"pension_inc{k}", f"{{P{k}}}'s eligible pension income", "INCOME TAX",
                 (lambda k: lambda r, i: f"={c(f'pension{k}', r)}+IF({c(f'age{k}', r)}>=65,{c(f'rrif{k}', r)},0)+{signed(k, r)}")(k), MONEY),
                (f"oas_clawback{k}", f"{{P{k}}}'s OAS clawback", "INCOME TAX",
                 (lambda k: lambda r, i: f"=MIN({c(f'oas{k}', r)},{TT['OAS_CLAWBACK_RATE']}*MAX(0,{c(f'net_income{k}', r)}-{TT['OAS_CLAWBACK_THRESHOLD']}*{TAXF(r)}))")(k), MONEY),
                (f"taxable_income{k}", f"{{P{k}}}'s taxable income", "INCOME TAX",
                 (lambda k: lambda r, i: f"=MAX(0,{c(f'net_income{k}', r)}-{c(f'oas_clawback{k}', r)})")(k), MONEY),
            ]
        cols.append(("medical_exp", "Medical expenses for the credit", "INCOME TAX",
                     lambda r, i: f"={c('health_total', r)}-{c('employer_health', r)}", MONEY, "Claimed by the spouse with the lower net income"))
        for k in (1, 2):
            claim = (lambda k: lambda r: (f"({c('net_income1', r)}<={c('net_income2', r)})" if k == 1 else
                                          f"({c('net_income1', r)}>{c('net_income2', r)})"))(k)
            ni = (lambda k: lambda r: c(f"net_income{k}", r))(k)
            age65 = (lambda k: lambda r: f"({c(f'age{k}', r)}>=65)")(k)
            cols += [
                (f"fed_credits{k}", f"{{P{k}}}'s federal credit amounts", "INCOME TAX",
                 (lambda k, claim, ni, age65: lambda r, i: (
                     f"=({TT['FED_BPA_MAX']}-({TT['FED_BPA_MAX']}-{TT['FED_BPA_MIN']})*MIN(1,MAX(0,({ni(r)}-{TT['FED_BPA_PHASE_START']}*{TAXF(r)})"
                     f"/(({TT['FED_BPA_PHASE_END']}-{TT['FED_BPA_PHASE_START']})*{TAXF(r)}))))*{TAXF(r)}"
                     f"+{age65(r)}*MAX(0,{TT['FED_AGE_AMOUNT']}*{TAXF(r)}-{TT['AGE_REDUCTION_RATE']}*MAX(0,{ni(r)}-{TT['FED_AGE_THRESHOLD']}*{TAXF(r)}))"
                     f"+MIN({TT['FED_PENSION_AMOUNT']},{c(f'pension_inc{k}', r)})"
                     f"+MIN({TT['FED_EMPLOYMENT_AMOUNT']}*{TAXF(r)},{c(f'pay{k}', r)})+{c(f'payroll{k}', r)}"
                     f"+{claim(r)}*MAX(0,{c('medical_exp', r)}-MIN(0.03*{ni(r)},{TT['FED_MEDICAL_THRESHOLD']}*{TAXF(r)}))"))(k, claim, ni, age65), MONEY,
                 "Basic personal, age, pension, employment, CPP/EI and medical amounts"),
                (f"fed_tax{k}", f"{{P{k}}}'s federal tax", "INCOME TAX",
                 (lambda k: lambda r, i: (f"=MAX(0," + prog(c(f'taxable_income{k}', r), "fed", TAXF(r))
                                          + f"-{TT['FED_CREDIT_RATE']}*{c(f'fed_credits{k}', r)}-{TT['FED_ELIGIBLE_DTC']}*{c(f'gross_div{k}', r)})"))(k), MONEY),
                (f"prov_credits{k}", f"{{P{k}}}'s provincial credit amounts", "INCOME TAX",
                 (lambda k, claim, ni, age65: lambda r, i: (
                     f"={TT['PROV_BPA']}*{PF(r)}"
                     f"+{age65(r)}*MAX(0,{TT['PROV_AGE_AMOUNT']}*{PF(r)}-{TT['AGE_REDUCTION_RATE']}*MAX(0,{ni(r)}-{TT['PROV_AGE_THRESHOLD']}*{PF(r)}))"
                     f"+MIN({TT['PROV_PENSION_AMOUNT']},{c(f'pension_inc{k}', r)})+{c(f'payroll{k}', r)}"
                     f"+{claim(r)}*MAX(0,{c('medical_exp', r)}-MIN(0.03*{ni(r)},{TT['PROV_MEDICAL_THRESHOLD']}*{PF(r)}))"))(k, claim, ni, age65), MONEY),
                (f"prov_tax{k}", f"{{P{k}}}'s provincial tax", "INCOME TAX",
                 (lambda k, ni: lambda r, i: (f"=MAX(0," + prog(c(f'taxable_income{k}', r), "prov", PF(r))
                                              + f"-{TT['PROV_CREDIT_RATE']}*{c(f'prov_credits{k}', r)}-{TT['PROV_ELIGIBLE_DTC']}*{c(f'gross_div{k}', r)}"
                                              + f"-MAX(0,{TT['PROV_REDUCTION_MAX']}*{PF(r)}-{TT['PROV_REDUCTION_RATE']}*MAX(0,{ni(r)}-{TT['PROV_REDUCTION_THRESHOLD']}*{PF(r)})))"))(k, ni), MONEY),
            ]
        cols.append(("income_tax", "Income tax for the year (incl. OAS clawback)", "INCOME TAX",
                     lambda r, i: (f"={c('fed_tax1', r)}+{c('fed_tax2', r)}+{c('prov_tax1', r)}+{c('prov_tax2', r)}"
                                   f"+{c('oas_clawback1', r)}+{c('oas_clawback2', r)}"), MONEY))
        return cols

    def results_block(self):
        ws, L, first, last, H = self.ws, self.letters, self.first, self.last, self.H
        yr = f"${L['year']}${first}:${L['year']}${last}"
        col = lambda k: f"${L[k]}${first}:${L[k]}${last}"     # noqa: E731
        c0 = 5
        ws.cell(4, c0, "RESULTS (today's dollars)").font = WHITE_BOLD
        for cc in range(c0, c0 + 3):
            ws.cell(4, cc).fill = BAND
        for i, (key, text, kind, age) in enumerate(RESULTS):
            r = RESULTS_ROW0 + i
            ws.cell(r, c0, self.label(text))
            if kind == "nw":
                f = f'=IFERROR(INDEX({col("net_worth_today")},MATCH({H["p1_born"]}+{age},{yr},0)),"")'
            elif kind == "inv":
                f = f'=IFERROR(MAX(0,INDEX({col("investments_today")},MATCH({H["p1_born"]}+{age},{yr},0))),"")'
            elif kind == "short":
                f = f'=IF(COUNTIF({col("shortfall")},">1")=0,"never",SUMPRODUCT(MIN({yr}+({col("shortfall")}<=1)*10000)))'
            elif kind == "living1":
                f = f"=INDEX({col('living_costs')},1)"
            elif kind == "rate1":
                f = f"=INDEX({col('draw_rate')},1)"
            else:
                f = f"=SUMPRODUCT({col('income_tax')}/{col('infl_f')})"
            cell = ws.cell(r, c0 + 2, f)
            cell.number_format = PCT if kind == "rate1" else (YEAR if kind == "short" else MONEY)
            cell.fill, cell.font = KEY, BOLD
        ws.column_dimensions["E"].width = 40
        ws.column_dimensions["G"].width = 14

    def readme(self):
        ws = self.wb.create_sheet("README", 0)
        us = self.us
        lines = [
            ("Retirement plan", TITLE),
            ("A year-by-year retirement model for a couple, with a Monte Carlo test against history. See the project README for how to use it with Claude.", None),
            ("", None),
            ("Tabs", BOLD),
            ("Model - ...: a scenario. Yellow cells are its inputs (when you stop working, pensions, property sales, one-off events). Duplicate the tab to try another scenario.", None),
            ("household: names, birth years, " + ("Social Security benefits" if us else "CPP and OAS") + ", kids.", None),
            ("budget: 12 months of spending by category (the statement pipeline fills it). The retirement_budget column is what the model spends.", None),
            ("kids_budget: one line per kid and cost (support, college, help after college), each with its own years and amount.", None),
            ("other_budget: renovations, cars and other big irregular costs, per year.", None),
            ("properties: your home and any other properties.", None),
            ("healthcare: " + ("marketplace, Medicare and long-term care costs." if us else "private health/dental plan, out-of-pocket and long-term care costs."), None),
            ("assets: every account, its balance and which bucket it feeds.", None),
            ("tax_tables: the tax law the model uses (" + ("federal + your state" if us else "federal + your province") + "). Edit to change state/province or update the law.", None),
            ("monte_carlo: settings and results of the Monte Carlo, which compares every Model tab (run it from the command line).", None),
            ("historical_returns, life_table: data for the Monte Carlo.", None),
            ("", None),
            ("What the model does each year", BOLD),
            ("1. Adds up living costs: household budget, housing, kids, college shortfall, other costs, healthcare, long-term care, one-off expenses.", None),
            ("2. Adds up income: take-home pay, " + ("Social Security" if us else "CPP and OAS") + ", property income, dividends, required withdrawals, sales, one-off income.", None),
            ("3. Covers any gap from " + ("taxable accounts, then pre-tax (401k/IRA), then Roth, then HSA" if us else "non-registered accounts, then RRSP/RRIF, then TFSA") + "; a surplus is reinvested.", None),
            ("4. Computes that year's income tax on what actually happened and pays it the following April.", None),
            ("", None),
            ("Colors: yellow = input you can change; green = headline result; blue bands group the model columns.", None),
            ("Not financial advice. Tax rules are simplified; check important decisions with a professional.", None),
        ]
        for i, (text, font) in enumerate(lines):
            c = ws.cell(i + 1, 1, text)
            if font:
                c.font = font
        ws.column_dimensions["A"].width = 150

    def build(self, out):
        self.household()
        self.budget()
        self.kids_budget()
        self.other_budget()
        self.properties()
        self.healthcare()
        self.assets()
        self.tax_tables()
        self.data_tabs()
        self.monte_carlo_tab()
        self.model()
        self.readme()
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(out)
        return out


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    inputs, out = sys.argv[1], sys.argv[2]
    if Path(out).exists():
        sys.exit(f"{out} already exists - not overwriting your plan. Pick another name or delete it first.")
    Book(inputs).build(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
