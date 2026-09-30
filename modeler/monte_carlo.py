#!/usr/bin/env python3
"""Test every scenario against history: rerun the model once per historical starting year, and more.

Reads the plan workbook - each 'Model - ...' tab's yellow inputs plus the shared tabs - and runs model.py with each
year's investment return and inflation taken from history instead of the Model tab's flat rates. Starting year
1928 means the plan's first year gets 1928's returns and inflation, the second 1929's, and so on, wrapping around
at the end of the data, so every year of history is a starting point. Random runs stitch together randomly chosen
multi-year blocks of history.

Every market path is run twice: once with you both alive to the end of the table (the stress test), and once or
more with lifetimes drawn from the life_table tab (death rates scaled to add the extra years of life expectancy set
on the monte_carlo tab). In lifetime runs household and other spending drop to the survivor share after the first
death, the survivor keeps the larger government pension (plus a CPP survivor pension in Canada), the deceased's
pre-tax accounts roll over to the survivor, and long-term care is each person's final years. "Money outlasts you
both" is the share of lifetime runs that never run short while either of you is alive.

Before simulating, each scenario is rerun at its own flat rates and compared with its Model tab's results, so a
misread input shows up as a MISMATCH instead of a silently wrong answer.

Results are written to the workbook's monte_carlo tab (a backup copy of the workbook is saved first) and printed.

Usage: monte_carlo.py <plan.xlsx> [--no-write]
"""
import argparse
import random
import shutil
import statistics
import sys
import time
from pathlib import Path

import model as M
from workbook import compare_headline, headline, num, read_plan

# monte_carlo tab inputs: (row, label, default, note, number format)
SETTINGS = [
    (5, "Stocks", 0.6, "Share of investments in stocks (historical_returns 'stocks' column)", "0%"),
    (6, "Bonds", 0.4, "Share in bonds ('bonds' column)", "0%"),
    (7, "Cash", 0.0, "Share in cash ('bills' column); the three should add to 100%", "0%"),
    (8, "Random runs per scenario", 1000, "0 = historical starting years only", None),
    (9, "Block length for random runs (years)", 10, "Random runs stitch together blocks of this many consecutive years of history", None),
    (10, "Random seed", 1, "Change it to draw a different set of random runs", None),
    (11, "Scenarios to run", "", "Comma-separated Model tab names; blank = every Model tab", None),
    (12, "{P1}: extra years of life expectancy", 3, "Years added to the life table's average. Healthier, higher-income people live longer "
         "(often +3 to +6). More years is the cautious direction", None),
    (13, "{P2}: extra years of life expectancy", 3, "Same for person 2", None),
    (14, "Survivor spending", 0.7, "After the first death, household and other spending as a share of the couple's", "0%"),
    (15, "Lifetimes drawn per historical start", 20, "Each historical start is run with this many random pairs of lifetimes", None),
]
RESULTS_ROW = 18


def settings_labels(names):
    return [label.replace("{P1}", names["P1"]).replace("{P2}", names["P2"]) for _, label, _, _, _ in SETTINGS]


def read_settings(ws):
    vals = []
    for row, _, default, _, _ in SETTINGS:
        v = ws.cell(row, 2).value if ws is not None else None
        vals.append(default if v in (None, "") and row != 11 else v)
    ws_, wb_, wc_, runs, block, seed, names, extra1, extra2, survivor, draws = vals
    return dict(mix=(num(ws_), num(wb_), num(wc_)), runs=int(num(runs)), block=max(1, int(num(block))), seed=int(num(seed)),
                names=str(names or ""), extra=(num(extra1), num(extra2)), survivor=num(survivor), draws=max(1, int(num(draws))))


# ---------------------------------------------------------------------------------------------- simulation

def history_rows(ws):
    head = [c.value for c in ws[1]]
    idx = {h: head.index(h) for h in ("year", "stocks", "bonds", "bills", "inflation")}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["year"]] in (None, ""):
            break
        out.append({k: (int(row[i]) if k == "year" else float(row[i])) for k, i in idx.items()})
    return out


def paths(hist, years, mix, runs, block, seed):
    """(kind, label, returns by year, inflation by year) for every historical start, then the random runs."""
    ws, wb, wc = mix
    n = len(hist)

    def build(idx):
        ret = {y: ws * hist[i]["stocks"] + wb * hist[i]["bonds"] + wc * hist[i]["bills"] for y, i in zip(years, idx)}
        return ret, {y: hist[i]["inflation"] for y, i in zip(years, idx)}

    out = [("history", hist[s]["year"], *build([(s + k) % n for k in range(len(years))])) for s in range(n)]
    rng = random.Random(seed)
    for r in range(runs):
        idx = []
        while len(idx) < len(years):
            s = rng.randrange(n)
            idx += [(s + k) % n for k in range(block)]
        out.append(("random", r + 1, *build(idx[:len(years)])))
    return out


def life_table(ws):
    """{age: (male, female) probability of dying within a year}"""
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            break
        out[int(row[0])] = (float(row[1]), float(row[2]))
    return out


def survival(q, age0, mult=1.0):
    alive, out = 1.0, {}
    for a in range(age0, max(q) + 2):
        out[a] = alive
        alive *= 1 - min(1.0, q.get(a, 1.0) * mult)
    return out


def life_exp(surv):
    return sum(surv.values()) - 0.5


def adjusted_survival(q, age0, extra):
    """Survival with life expectancy `extra` years above the table's: every death rate scaled by one factor."""
    target = life_exp(survival(q, age0)) + extra
    lo, hi = 0.01, 5.0
    for _ in range(60):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if life_exp(survival(q, age0, mid)) > target else (lo, mid)
    return survival(q, age0, (lo + hi) / 2)


def death_year(surv, born, rng, last):
    """First year this person is no longer alive; None = alive past the end of the table."""
    u = rng.random()
    for age in sorted(surv):
        if surv.get(age + 1, 0.0) < u:
            return born + age + 1 if born + age + 1 <= last else None
    return None


def pct(xs, p):
    xs = sorted(xs)
    return xs[min(len(xs) - 1, int(p * len(xs)))]


SUMMARY_HEAD = ["Scenario", "Check vs Model tab (flat rates)", "Money outlasts you both: historical starts",
                "... random runs", "Investments left at the end: 10th percentile", "... median", "... 90th percentile",
                "Net worth left at the end (incl. property): median", "Both alive to the end - money lasts: historical starts",
                "... random runs", "... historical starts that run short", "... earliest shortfall",
                "Model tab (flat rates): investments at {P1} 90", "Model tab: first shortfall"]


def run_all(plan, settings):
    data0 = plan["data"]
    models = list(plan["models"])
    if settings["names"].strip():
        wanted = [n.strip() for n in settings["names"].split(",") if n.strip()]
        missing = [n for n in wanted if n not in models]
        if missing:
            sys.exit(f"not a Model tab: {', '.join(missing)}")
        models = wanted
    if abs(sum(settings["mix"]) - 1) > 0.001:
        print(f"warning: stock/bond/cash shares add to {sum(settings['mix']):.0%}, not 100%")
    wb = plan["wb"]
    hist = history_rows(wb["historical_returns"])
    people = data0["people"]
    summary, per_start, info = [], {}, ""
    for name in models:
        t0 = time.time()
        S, data, _, sheet = plan["models"][name]
        years = list(range(data["start"], data["end"] + 1))
        born1 = people[0]["born"]
        at90 = min(born1 + 90, data["end"])
        base = M.run(data, S)
        mismatch = compare_headline(sheet, headline(base, data))
        all_paths = paths(hist, years, settings["mix"], settings["runs"], settings["block"], settings["seed"])
        # lifetimes: the same draws for every scenario (same seed)
        q = life_table(wb["life_table"])
        surv = [adjusted_survival({a: v[people[i]["sex"]] for a, v in q.items()}, data["start"] - people[i]["born"], settings["extra"][i])
                for i in (0, 1)]
        rng = random.Random(settings["seed"])
        lifetimes = []
        for kind, _, ret, infl in all_paths:
            for _ in range(settings["draws"] if kind == "history" else 1):
                lifetimes.append((kind, ret, infl, death_year(surv[0], people[0]["born"], rng, data["end"]),
                                  death_year(surv[1], people[1]["born"], rng, data["end"])))
        if not info:
            either = lambda y: 1 - (1 - surv[0].get(y - people[0]["born"], 0)) * (1 - surv[1].get(y - people[1]["born"], 0))  # noqa: E731
            info = (f"life expectancy used: {people[0]['name']} to {data['start'] - people[0]['born'] + life_exp(surv[0]):.0f}, "
                    f"{people[1]['name']} to {data['start'] - people[1]['born'] + life_exp(surv[1]):.0f}; chance either of you "
                    f"is alive when {people[0]['name']} would be 90: {either(born1 + 90):.0%}, 95: {either(born1 + 95):.0%}, "
                    f"100: {either(born1 + 100):.0%}; survivor spending {settings['survivor']:.0%}")
        flat = {k: v for k, v in S.items() if k != "returns"}
        results = []
        for kind, label, ret, infl in all_paths:
            res = M.run(data, flat, returns=ret, inflation_path=infl)
            by = {r["year"]: r for r in res}
            short = next((r["year"] for r in res if r["shortfall"] > 1), None)
            results.append((kind, label, dict(short=short, inv90=by[at90]["investments_today"])))
        lives = []
        for kind, ret, infl, d1, d2 in lifetimes:
            res = M.run(data, flat, returns=ret, inflation_path=infl, p1_death_year=d1, p2_death_year=d2,
                        survivor_spending=settings["survivor"])
            lives.append((kind, dict(fails=any(r["shortfall"] > 1 for r in res), inv_end=res[-1]["investments_today"],
                                     nw_end=res[-1]["net_worth_today"])))
        summary.append(summarize(name, results, lives, headline(base, data), base, at90, mismatch))
        per_start[name] = [o for kind, _, o in results if kind == "history"]
        print(f"{name}: {len(results) + len(lives) + 1} runs in {time.time() - t0:.1f}s"
              + ("" if not mismatch else "  ** MISMATCH vs Model tab: " + "; ".join(mismatch)))
    return dict(hist=hist, models=models, summary=summary, per_start=per_start, info=info, settings=settings,
                years=len(years), people=people)


def summarize(name, runs, lives, head, base, at90, mismatch):
    hist = [o for kind, _, o in runs if kind == "history"]
    rand = [o for kind, _, o in runs if kind == "random"]
    fails = [(label, o["short"]) for kind, label, o in runs if kind == "history" and o["short"]]
    worst = min(fails, key=lambda f: f[1]) if fails else None
    lh = [o for kind, o in lives if kind == "history"]
    lr = [o for kind, o in lives if kind == "random"]
    ok = lambda os: sum(not o["fails"] for o in os) / len(os) if os else ""  # noqa: E731
    inv = [o["inv_end"] for o in lh]
    by = {r["year"]: r for r in base}
    return [name, "OK" if not mismatch else "MISMATCH - " + "; ".join(mismatch), ok(lh), ok(lr),
            pct(inv, 0.10) if inv else "", statistics.median(inv) if inv else "", pct(inv, 0.90) if inv else "",
            statistics.median(o["nw_end"] for o in lh) if lh else "",
            sum(o["short"] is None for o in hist) / len(hist), (sum(o["short"] is None for o in rand) / len(rand)) if rand else "",
            len(fails), f"{worst[0]} start runs short in {worst[1]}" if worst else "none",
            by[at90]["investments_today"], head["short"]]


def print_results(r):
    f = lambda v: f"{v:.0%}" if v != "" else "-"  # noqa: E731
    d = lambda v: f"${v:,.0f}" if v != "" else "-"  # noqa: E731
    print(r["info"])
    for (name, chk, out_h, out_r, inv10, inv50, inv90, nw50, last_h, last_r, _, worst, flat_inv, flat_short) in r["summary"]:
        print(f"\n{name}  [{chk}]")
        print(f"  money outlasts you both: {f(out_h)} of historical lifetime runs, {f(out_r)} of random ones")
        print(f"  left at the end: investments {d(inv10)} / {d(inv50)} / {d(inv90)} (10th/median/90th), net worth median {d(nw50)}")
        print(f"  both alive to the end: lasts in {f(last_h)} of historical starts, {f(last_r)} of random runs; earliest "
              f"shortfall {worst} (flat rates: {d(flat_inv)} investments at {r['people'][0]['name']} 90, shortfall {flat_short})")


def lock_files(path):
    p = Path(path)
    return [x for x in (p.with_name(f".~lock.{p.name}#"), p.with_name(f"~${p.name}")) if x.exists()]


def write_results(path, r):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    locks = lock_files(path)
    if locks:
        sys.exit(f"{path} looks open in Excel/LibreOffice ({locks[0].name}); close it and rerun (results printed above)")
    backup_dir = Path(path).parent / "backups"
    backup_dir.mkdir(exist_ok=True)
    backup = backup_dir / f"{Path(path).stem} {time.strftime('%Y-%m-%d %H%M%S')}.xlsx"
    shutil.copy2(path, backup)
    wb = openpyxl.load_workbook(path)
    ws = wb["monte_carlo"]
    for row in ws.iter_rows(min_row=RESULTS_ROW - 1, max_row=max(ws.max_row, RESULTS_ROW)):
        for c in row:
            c.value = None
            c.font, c.fill, c.number_format = Font(), PatternFill(), "General"
    s = r["settings"]
    mix = " / ".join(f"{w:.0%}" for w in s["mix"])
    bold = Font(bold=True)
    ws.cell(RESULTS_ROW - 1, 1, f"Last run {time.strftime('%Y-%m-%d %H:%M')}: stocks/bonds/cash {mix}; {len(r['hist'])} historical "
                                f"starts ({r['hist'][0]['year']}-{r['hist'][-1]['year']}), {s['runs']} random runs of "
                                f"{s['block']}-year blocks (seed {s['seed']}).")
    ws.cell(RESULTS_ROW, 1, r["info"][0].upper() + r["info"][1:] + ".")
    ws.cell(RESULTS_ROW + 1, 1, "RESULTS").font = bold
    p1 = r["people"][0]["name"]
    for j, h in enumerate(SUMMARY_HEAD):
        c = ws.cell(RESULTS_ROW + 2, j + 1, h.replace("{P1}", p1))
        c.font, c.alignment = bold, Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(r["summary"]):
        for j, v in enumerate(row):
            c = ws.cell(RESULTS_ROW + 3 + i, j + 1, v)
            if j in (2, 3, 8, 9):
                c.number_format = "0%"
            if j in (4, 5, 6, 7, 12):
                c.number_format = "$#,##0"
    start = RESULTS_ROW + 5 + len(r["summary"])
    ws.cell(start, 1, "BY HISTORICAL STARTING YEAR").font = bold
    head = ["Start year", "History used"] + [f"{m}: first shortfall" for m in r["models"]] + \
           [f"{m}: investments at {p1} 90" for m in r["models"]]
    for j, h in enumerate(head):
        c = ws.cell(start + 1, j + 1, h)
        c.font, c.alignment = bold, Alignment(wrap_text=True, vertical="top")
    n = len(r["hist"])
    for s_, h in enumerate(r["hist"]):
        end = r["hist"][(s_ + r["years"] - 1) % n]["year"]
        used = f"{h['year']}-{end}" + (" (wraps)" if s_ + r["years"] > n else "")
        row = [h["year"], used] + [r["per_start"][m][s_]["short"] or "never" for m in r["models"]] + \
              [r["per_start"][m][s_]["inv90"] for m in r["models"]]
        for j, v in enumerate(row):
            c = ws.cell(start + 2 + s_, j + 1, v)
            if j >= 2 + len(r["models"]):
                c.number_format = "$#,##0"
    wb.save(path)
    print(f"\nwrote the monte_carlo tab of {path} (backup: {backup})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--no-write", action="store_true", help="print results only; don't change the workbook")
    a = ap.parse_args()
    plan = read_plan(a.workbook)
    settings = read_settings(plan["wb"]["monte_carlo"] if "monte_carlo" in plan["wb"].sheetnames else None)
    r = run_all(plan, settings)
    print_results(r)
    if not a.no_write:
        write_results(a.workbook, r)


if __name__ == "__main__":
    main()
