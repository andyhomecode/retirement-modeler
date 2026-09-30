#!/usr/bin/env python3
"""One entry point for everything (the ./plan script runs this in its Python environment).

  new <us|ca> <folder>        start a plan: copies the example inputs and starter rules, builds <folder>/plan.xlsx
  check <folder>              compare the workbook's formulas with the Python model, cell by cell
  montecarlo <folder>         replay every scenario through history; results go to the monte_carlo tab
  statements <folder> [...]   import statements -> budget and balances (add --update-workbook to merge them in)
  inputs <folder> [tab]       list a Model tab's input keys, cells and values
  set <folder> <tab> key=value ...   change inputs (backup first); prints old -> new for the log
  duplicate <folder> <tab> <new name>  copy a Model tab to start a new scenario
  whatif <folder> [--tab T] --set key=value ... [--vs --set ...] [--mc]   try changes without saving them
  build <folder>              (re)create plan.xlsx from <folder>/inputs - only when plan.xlsx doesn't exist yet
  demo <us|ca>                copy the example to demo-<us|ca>/ and run statements, check and Monte Carlo on it

<folder> is a plan folder (plan.xlsx, inputs/, rules/, statements/, ledger/, PLANNING.md); a path to an .xlsx works
for check and montecarlo too.
"""
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def workbook(arg):
    p = Path(arg)
    return p if p.suffix.lower() == ".xlsx" else p / "plan.xlsx"


def new(country, folder):
    country = country.lower()
    src = ROOT / "examples" / country
    if not src.exists():
        sys.exit("country must be us or ca")
    dest = Path(folder)
    if (dest / "plan.xlsx").exists():
        sys.exit(f"{dest / 'plan.xlsx'} already exists")
    dest.mkdir(parents=True, exist_ok=True)
    shutil.copytree(src / "inputs", dest / "inputs", dirs_exist_ok=True)
    shutil.copytree(ROOT / "templates" / "rules", dest / "rules", dirs_exist_ok=True)
    (dest / "statements" / "converted").mkdir(parents=True, exist_ok=True)
    (dest / "ledger").mkdir(exist_ok=True)
    readme = dest / "statements" / "README.txt"
    if not readme.exists():
        shutil.copy(ROOT / "templates" / "statements_README.txt", readme)
    if not (dest / "PLANNING.md").exists():
        shutil.copy(ROOT / "templates" / "PLANNING.md", dest / "PLANNING.md")
    if not (dest / "CLAUDE.md").exists():
        text = (ROOT / "templates" / "CLAUDE.md").read_text().replace("{REPO}", str(ROOT))
        (dest / "CLAUDE.md").write_text(text)
    build(dest)
    print(f"\nNew plan in {dest}. It starts as a copy of the {country.upper()} example household - replace the inputs with "
          f"yours (or ask Claude to). See HOW_TO_USE_WITH_CLAUDE.md.")


def build(folder):
    from build_workbook import Book
    folder = Path(folder)
    out = folder / "plan.xlsx"
    if out.exists():
        sys.exit(f"{out} already exists - it's your plan now; edit it rather than rebuilding (delete it first if you really mean to)")
    Book(folder / "inputs").build(out)
    print(f"wrote {out}")


def demo(country):
    from check_workbook import check
    import monte_carlo
    import pipeline
    country = country.lower()
    dest = ROOT / f"demo-{country}"
    if dest.exists():
        shutil.rmtree(dest)
    shutil.copytree(ROOT / "examples" / country, dest, ignore=shutil.ignore_patterns("ledger", "backups"))
    print(f"== copied examples/{country} to {dest.name}/\n\n== statements")
    pipeline.main([str(dest), "--update-workbook"])
    print("\n== check")
    check(dest / "plan.xlsx")
    print("\n== Monte Carlo")
    sys.argv = ["monte_carlo.py", str(dest / "plan.xlsx")]
    monte_carlo.main()


def list_inputs(path, tab=None):
    import openpyxl
    from build_workbook import input_spec
    from workbook import input_cells, kv
    wb = openpyxl.load_workbook(path)
    hh = kv(wb["household"])
    country = str(hh["country"]).strip().upper()
    labels = {item[0]: item[1].replace("{P1}", str(hh["p1_name"])).replace("{P2}", str(hh["p2_name"]))
              for item in input_spec(country) if item[0] != "section"}
    calc = {item[0] for item in input_spec(country) if item[0] != "section" and item[5] == "calc"}
    tabs = [tab] if tab else [t for t in wb.sheetnames if t.startswith("Model - ")]
    cells = input_cells(country)
    for t in tabs:
        ws = wb[t]
        print(f"== {t}")
        for k, c in cells.items():
            v = ws[c].value
            if k.startswith(("prop", "event")) and v in (None, ""):
                continue
            shown = "(calculated)" if k in calc or (isinstance(v, str) and v.startswith("=") and k.startswith("prop")) else \
                ("" if v is None else str(v))
            print(f"  {k:<20} {c:<5} {shown[:28]:<28} {labels.get(k, '')}")


def main():
    a = sys.argv[1:]
    if not a or a[0] in ("-h", "--help", "help"):
        sys.exit(__doc__)
    cmd, rest = a[0], a[1:]
    if cmd == "new" and len(rest) == 2:
        new(*rest)
    elif cmd == "build" and len(rest) == 1:
        build(rest[0])
    elif cmd == "check" and rest:
        from check_workbook import check
        sys.exit(0 if check(workbook(rest[0]), "--verbose" in rest) else 1)
    elif cmd == "montecarlo" and rest:
        import monte_carlo
        sys.argv = ["monte_carlo.py", str(workbook(rest[0]))] + rest[1:]
        monte_carlo.main()
    elif cmd == "statements" and rest:
        import pipeline
        pipeline.main(rest)
    elif cmd == "inputs" and rest:
        list_inputs(workbook(rest[0]), rest[1] if len(rest) > 1 else None)
    elif cmd == "set" and len(rest) >= 3:
        from workbook import parse_value, set_inputs
        changes = {}
        for kv_ in rest[2:]:
            k, _, v = kv_.partition("=")
            changes[k.strip()] = parse_value(v)
        done, dest = set_inputs(workbook(rest[0]), rest[1], changes)
        for k, cell, old, new_ in done:
            print(f"{rest[1]}!{cell} {k}: {old!r} -> {new_!r}")
        print(f"backup: {dest}\nNow run ./plan check {rest[0]} (and note the change in PLANNING.md)")
    elif cmd == "duplicate" and len(rest) == 3:
        from workbook import duplicate_tab
        name, dest = duplicate_tab(workbook(rest[0]), rest[1], rest[2])
        print(f"added {name!r} (backup: {dest})")
    elif cmd == "whatif" and rest:
        import whatif
        whatif.main([str(workbook(rest[0]))] + rest[1:])
    elif cmd == "demo" and len(rest) == 1:
        demo(rest[0])
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
